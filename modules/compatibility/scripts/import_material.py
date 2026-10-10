#!/usr/bin/env python3
"""Import existing QA material without guessing missing business rules."""

from __future__ import annotations

import argparse
import copy
import csv
from datetime import date, datetime, time, timezone
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile
from urllib.parse import urlsplit

MODULES = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(MODULES / "testpoints" / "scripts"))
from point_common import validate_points
from qa_cases import (HEADERS, METHODS, SCOPES, ValidationError, validate_cases,
                      validate_spec, write_text)
sys.path.insert(0, str(MODULES / "req-processor" / "scripts"))
from validate_spec import check as check_spec

TARGETS = {"requirements": "REQ清单.json", "testpoints": "测试点数据.json", "cases": "用例数据.json"}
POINT_FIELDS = ["测试点编号", "优先级", "测试点名称", "前置条件", "检查内容", "测试数据",
                "预期结果", "关联需求", "验收标准", "场景类型", "设计方法"]
ALIASES = {
    "id": ["id", "requirement_id", "需求编号", "REQ编号"],
    "title": ["title", "requirement_name", "name", "需求名称"],
    "description": ["description", "需求描述", "描述"],
    "detailed_description": ["detailed_description", "完整描述", "详细描述"],
    "kind": ["kind", "type", "需求类型"],
    "status": ["status", "需求状态", "状态"],
    "acceptance_criteria": ["acceptance_criteria", "验收条件", "验收标准"],
    "测试点编号": ["测试点编号", "测试点ID", "testpoint_id", "test_point_id", "point_id", "id"],
    "测试点名称": ["测试点名称", "测试点", "testpoint_name", "test_point_name", "title", "name"],
    "用例编号": ["用例编号", "用例ID", "case_id", "testcase_id", "id"],
    "用例名称": ["用例名称", "用例标题", "case_name", "testcase_name", "title", "name"],
    "优先级": ["优先级", "priority"],
    "前置条件": ["前置条件", "前提条件", "preconditions", "precondition"],
    "检查内容": ["检查内容", "检查点", "test_steps", "test_content", "checks", "操作步骤"],
    "操作步骤": ["操作步骤", "测试步骤", "test_steps", "steps"],
    "测试数据": ["测试数据", "test_data", "data"],
    "预期结果": ["预期结果", "期望结果", "expected_results", "expected_result", "expected"],
    "实际结果": ["实际结果", "actual_result", "actual_results"],
    "关联需求": ["关联需求", "关联REQ", "requirement_id", "requirement_ids", "related_requirements", "requirements"],
    "验收标准": ["验收标准", "验收标准编号", "mapped_criteria", "acceptance_criteria", "criteria_ids"],
    "场景类型": ["场景类型", "scope", "scenario_type"],
    "设计方法": ["设计方法", "test_method", "design_methods", "methods"],
    "是否可自动化": ["是否可自动化", "可自动化", "automatable", "automation"],
    "自动化脚本": ["自动化脚本", "automation_script", "script"],
    "自动化建议": ["自动化建议", "自动化建议分类", "automation_advice", "advice"],
    "自动化建议明细": ["自动化建议明细", "建议明细", "automation_advice_detail", "advice_detail"],
    "关联测试点": ["关联测试点", "测试点", "testpoint_ids", "related_testpoints"],
}
KINDS = {"functional": "functional", "constraint": "constraint", "nonfunctional": "nonfunctional",
         "功能需求": "functional", "功能需求(FR)": "functional", "FR": "functional",
         "非功能需求": "nonfunctional", "非功能需求(NFR)": "nonfunctional", "NFR": "nonfunctional",
         "业务规则": "constraint", "业务规则(BR)": "constraint", "BR": "constraint",
         "数据需求": "constraint", "数据需求(DR)": "constraint", "DR": "constraint"}
STATUSES = {"active": "active", "blocked": "blocked", "deprecated": "deprecated",
            "已确认": "active", "待澄清": "blocked", "阻塞": "blocked", "废弃": "deprecated"}
DERIVED = ["测试点.md", "测试点.xmind", "测试点覆盖率数据.json", "测试点覆盖率报告.md", "测试点覆盖矩阵.md",
           "测试点评审数据.json", "测试点评审结果.json", "测试点评审报告.md", "测试点用例关联数据.json", "测试点用例关联报告.md",
           "详细用例.xlsx", "详细用例_已评审.xlsx", "覆盖率数据.json", "需求覆盖率报告.md", "评审数据.json", "评审报告.md",
           "评审结果.json", "用例覆盖矩阵.md", "执行准备数据.json", "执行准备清单.md",
           "可读性数据.json", "可读性报告.md"]


def read_json(path):
    try:
        return json.loads(Path(path).read_text(encoding="utf-8-sig"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise ValidationError(f"无法读取JSON {Path(path).name}：{exc}") from exc


def write_json(path, value):
    def dates(value):
        if isinstance(value, (date, datetime, time)):
            return value.isoformat()
        raise TypeError(f"不支持JSON转换的单元格类型：{type(value).__name__}")
    write_text(Path(path), json.dumps(value, ensure_ascii=False, indent=2, default=dates) + "\n")


def token(value):
    return re.sub(r"[\s_\-（）()]", "", str(value)).casefold()


def text(value):
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(text(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def items(value):
    if value is None or value == "":
        return []
    if isinstance(value, list):
        return [text(item).strip() for item in value if text(item).strip()]
    if not isinstance(value, str):
        return [text(value)]
    return [part.strip() for part in re.split(r"[,，;；、\n]+", value) if part.strip()]


def steps(value):
    if isinstance(value, list):
        return [text(item) for item in value if text(item).strip()]
    return [str(value)] if value is not None and str(value).strip() else []


def md_literal(value):
    # Preserve exact original text in JSON and immutable evidence; escape image
    # instructions when rendering the imported description as a Markdown view.
    return text(value).replace("![", "!\\[").replace("<img", "&lt;img")


def canonical_fields(kind):
    if kind == "requirements":
        return ["id", "title", "description", "detailed_description", "kind", "status", "acceptance_criteria"]
    return POINT_FIELDS if kind == "testpoints" else HEADERS + ["验收标准", "场景类型", "设计方法", "关联测试点"]


def unresolved_markers(record):
    """Do not hide declared pending work inside a legacy copy during import."""
    problems = []
    if "draft" in record and record["draft"] is not False:
        problems.append("原资料draft未明确为false；需核对草稿后才可发布")
    if record.get("needs_redesign") or record.get("missing_fields"):
        problems.append("原资料仍有重新设计/缺失字段标记，需完成核对后才可发布")
    if record.get("update_state") is not None and record["update_state"] != "resolved":
        problems.append(f"原资料更新状态仍为{record['update_state']}，需明确resolved后才可发布")
    return problems


def field_map(keys, kind, mapping, *, table=False):
    """Ambiguous aliases need an explicit field mapping, never a first match."""
    result = {}
    overrides = mapping.get("fields", {})
    if not isinstance(overrides, dict):
        raise ValidationError("mapping.fields 必须为对象：规范字段 → 原字段名或1起始列号。")
    for field in canonical_fields(kind):
        if field in overrides:
            chosen = overrides[field]
            if isinstance(chosen, int) and not isinstance(chosen, bool) and table:
                if chosen < 1 or chosen > len(keys):
                    raise ValidationError(f"字段 {field} 的列号 {chosen} 越界。")
                result[field] = keys[chosen - 1]
            elif isinstance(chosen, str) and chosen in keys:
                result[field] = chosen
            else:
                raise ValidationError(f"字段 {field} 映射的原字段不存在：{chosen}。")
            continue
        aliases = {token(key) for key in ALIASES[field]}
        matches = [key for key in keys if token(key) in aliases]
        if len(matches) == 1:
            result[field] = matches[0]
        elif len(matches) > 1:
            raise ValidationError(f"字段 {field} 有多个候选列：{matches}；请通过 --mapping 指定。")
    return result


def convert_row(row, kind, mapping):
    fields = field_map(list(row), kind, mapping)
    defaults = mapping.get("defaults", {})
    if not isinstance(defaults, dict):
        raise ValidationError("mapping.defaults 必须为对象。")
    converted = {key: copy.deepcopy(defaults[key]) for key in canonical_fields(kind) if key in defaults}
    converted.update({field: row[source] for field, source in fields.items() if row[source] is not None and row[source] != ""})
    return converted


def json_rows(data, kind, module_name):
    rows = []
    if isinstance(data, list):
        rows = [(row, f"$[{index}]") for index, row in enumerate(data)]
    elif isinstance(data, dict):
        key = {"requirements": "requirements", "testpoints": "testpoints", "cases": "cases"}[kind]
        if isinstance(data.get(key), list):
            rows = [(row, f"$.{key}[{index}]") for index, row in enumerate(data[key])]
        elif kind == "requirements" and isinstance(data.get("modules"), list):
            modules = data["modules"]
            matched = [(i, item) for i, item in enumerate(modules) if item.get("module_name", item.get("name")) == module_name]
            if not matched:
                if len(modules) == 1:
                    matched = [(0, modules[0])]
                else:
                    raise ValidationError("旧JSON含多个模块；--module-dir 的目录名或 mapping.module 必须明确匹配 module_name。")
            if len(matched) != 1:
                raise ValidationError("旧JSON中模块名称重复，无法确定导入范围。")
            i, selected = matched[0]
            if not isinstance(selected.get("requirements"), list):
                raise ValidationError("旧模块 requirements 必须为数组。")
            rows = [(row, f"$.modules[{i}].requirements[{index}]") for index, row in enumerate(selected["requirements"])]
        elif kind == "testpoints" and isinstance(data.get("requirements"), list):
            for i, req in enumerate(data["requirements"]):
                if not isinstance(req, dict) or not isinstance(req.get("testpoints"), list):
                    raise ValidationError("旧测试点 requirements 每项应含 testpoints 数组。")
                for j, row in enumerate(req["testpoints"]):
                    merged = dict(row)
                    if "requirement_id" not in merged:
                        merged["requirement_id"] = req.get("requirement_id")
                    rows.append((merged, f"$.requirements[{i}].testpoints[{j}]"))
        elif kind == "requirements" and ("requirement_id" in data or "id" in data):
            rows = [(data, "$")]
    if not rows:
        raise ValidationError(f"没有找到可导入的{kind}记录；摘要索引不等于完整需求，需提供含描述和验收标准的明细。")
    if any(not isinstance(row, dict) for row, _ in rows):
        raise ValidationError("每条导入记录必须是对象。")
    return rows


def table_rows(path, kind, sheet, header_row, mapping):
    if path.suffix.lower() in (".csv", ".tsv"):
        try:
            with path.open(encoding="utf-8-sig", newline="") as stream:
                matrix = list(csv.reader(stream, delimiter="\t" if path.suffix.lower() == ".tsv" else ","))
        except UnicodeError as exc:
            raise ValidationError("表格需使用UTF-8编码；请另存为UTF-8 CSV/TSV后重试。") from exc
        label = path.name
    else:
        if path.suffix.lower() != ".xlsx":
            raise ValidationError("支持JSON、CSV、TSV、XLSX；旧.xls和含宏.xlsm请另存副本后导入。")
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValidationError("只读XLSX导入需要openpyxl；请使用工作区Python运行环境，或另存为CSV/TSV。") from exc
        try:
            book = load_workbook(path, read_only=True, data_only=False)
            if sheet is None and len(book.sheetnames) != 1:
                raise ValidationError("XLSX含多个工作表；请通过 --sheet 明确选择。")
            chosen = sheet or book.sheetnames[0]
            if chosen not in book.sheetnames:
                raise ValidationError(f"不存在工作表 {chosen}；可选：{book.sheetnames}。")
            matrix = [list(row) for row in book[chosen].iter_rows(values_only=True)]
            label = chosen
        finally:
            if "book" in locals():
                book.close()
    if header_row < 1 or header_row > len(matrix):
        raise ValidationError("--header-row 超出实际行数。")
    raw_headers = matrix[header_row - 1]
    while raw_headers and raw_headers[-1] is None:
        raw_headers.pop()
    headers = [text(value).strip() for value in raw_headers]
    duplicate = sorted({value for value in headers if value and headers.count(value) > 1})
    if duplicate:
        # Explicit numeric mappings permit duplicate user headings without data loss.
        if not mapping.get("fields") or any(not isinstance(v, int) for v in mapping["fields"].values()):
            raise ValidationError(f"表格标题重复：{duplicate}；请用 mapping.fields 的1起始列号映射所需字段。")
        headers = [f"@column{index}" for index in range(1, len(headers) + 1)]
    else:
        headers = [value or f"@column{index}" for index, value in enumerate(headers, 1)]
    resolved = field_map(headers, kind, mapping, table=True)
    rows = []
    for row_number, row in enumerate(matrix[header_row:], header_row + 1):
        if not any(value is not None and value != "" for value in row):
            continue
        if any(value is not None and value != "" for value in row[len(headers):]):
            raise ValidationError(f"{label} 第{row_number}行有无标题数据列；请补标题或调整 --header-row。")
        original = {key: row[index] if index < len(row) else None for index, key in enumerate(headers)}
        converted = {field: original[key] for field, key in resolved.items()}
        # Retain original headings and ignored columns alongside normalized fields.
        rows.append((converted, f"{label}:row:{row_number}", original))
    if not rows:
        raise ValidationError("表格没有非空数据行。")
    return rows


def make_id(old, prefix, code, registry, used):
    old = text(old).strip()
    if not old:
        return ""
    if old in used:
        raise ValidationError(f"重复原编号：{old}。")
    used.add(old)
    group = registry.setdefault(prefix, {})
    if old in group:
        new = group[old]
    elif re.fullmatch(rf"{prefix}_{code}_[0-9]{{3,}}", old):
        new = old
    else:
        occupied = {int(match.group(1)) for value in group.values()
                    if (match := re.fullmatch(rf"{prefix}_{code}_([0-9]{{3,}})", value))}
        number = 1
        while number in occupied:
            number += 1
        new = f"{prefix}_{code}_{number:03d}"
    if not re.fullmatch(rf"{prefix}_{code}_[0-9]{{3,}}", new) or new.endswith("_000"):
        raise ValidationError(f"编号映射无效：{old} → {new}。")
    if new in group.values() and group.get(old) != new:
        raise ValidationError(f"多个原编号映射到同一编号 {new}；请核对 mapping.ids。")
    group[old] = new
    return new


def resolve_links(value, prefix, registry):
    # 与 resolve_criteria 同一逃生通道：旧编号本身可能含 ，；、，映射表里整串登记的
    # 条目必须优先于按分隔符拆分，否则显式 mapping 对 REQ/测试点静默失效。
    whole = value.strip() if isinstance(value, str) else None
    candidates = [whole] if whole and whole in registry.get(prefix, {}) else items(value)
    return [registry.get(prefix, {}).get(item, item) for item in candidates]


def resolve_criteria(value, related, spec, registry):
    result, problems = [], []
    owners = {ac["id"]: req["id"] for req in spec["requirements"] for ac in req["acceptance_criteria"]}
    choices = [ac for req in spec["requirements"] if req["id"] in related for ac in req["acceptance_criteria"]]
    # A legacy scalar may be a complete acceptance sentence, not an ID list.
    # Preserve its punctuation for exact matching before trying delimiters.
    whole = value.strip() if isinstance(value, str) else None
    if whole and (whole in registry.get("AC", {}) or whole in owners
                  or any(ac["text"].strip() == whole for ac in choices)):
        candidates = [whole]
    else:
        candidates = items(value)
    for old in candidates:
        mapped = registry.get("AC", {}).get(old, old)
        if mapped in owners:
            if owners[mapped] not in related:
                problems.append(f"验收标准 {mapped} 不属于关联需求")
            else:
                result.append(mapped)
            continue
        matches = [ac["id"] for ac in choices if ac["text"].strip() == old.strip()]
        if len(matches) == 1:
            result.append(matches[0])
        else:
            problems.append(f"验收标准待绑定：{old}（文本精确匹配{'存在歧义' if matches else '未找到'}）")
    if not result:
        problems.append("验收标准未明确绑定，不能用该REQ的全部AC代填")
    return result, problems


def legacy_assets(original, source, batch, mapping):
    result, missing = [], []
    assets = original.get("related_images", original.get("import_images", []))
    if isinstance(assets, str):
        if assets.lstrip().startswith("["):
            try:
                assets = json.loads(assets)
            except json.JSONDecodeError:
                return [], ["related_images 的JSON文本无法解析"]
        else:
            assets = items(assets)
    if not isinstance(assets, list):
        return [], ["related_images 应为数组"]
    authorised = mapping.get("images", {})
    for asset in assets:
        descriptor = asset if isinstance(asset, dict) else {"path": asset}
        name = descriptor.get("path", descriptor.get("filename"))
        if not isinstance(name, str) or not name:
            missing.append("图片缺少本地路径")
            continue
        raw = authorised.get(name, name)
        asset_name = descriptor.get("import_asset")
        if isinstance(asset_name, str) and Path(asset_name).name == asset_name and name not in authorised:
            candidate = source.parent / "assets" / asset_name
            if not (source.parent / raw).is_file() and candidate.is_file():
                raw = f"assets/{asset_name}"
        parsed = urlsplit(raw)
        path = Path(raw)
        if parsed.scheme or parsed.netloc or parsed.query or parsed.fragment:
            missing.append(f"图片必须是授权的本地文件：{name}")
            continue
        if path.is_absolute() and name not in authorised:
            missing.append(f"绝对图片路径需 mapping.images 明确授权：{name}")
            continue
        path = path.resolve() if path.is_absolute() else (source.parent / path).resolve()
        if name not in authorised and not path.is_relative_to(source.parent.resolve()):
            missing.append(f"图片越出源文件目录，需 mapping.images 明确授权：{name}")
            continue
        if not path.is_file():
            missing.append(f"关联图片不存在：{name}")
            continue
        digest = hashlib.sha256(path.read_bytes()).hexdigest()[:12]
        destination = batch / "assets" / f"{digest}-{path.name}"
        destination.parent.mkdir(exist_ok=True)
        if not destination.exists():
            shutil.copy2(path, destination)
        result.append({"path": f"images/{destination.name}", "description": text(descriptor.get("description")),
                       "transcription": text(descriptor.get("transcription")), "import_asset": destination.name})
    return result, missing


def spec_markdown(data):
    lines = [f"# {data['module']}需求规范", "", "## 0. 基本信息", "",
             f"模块：{data['module']}；缩写：{data['code']}；导入版本：1.0。",
             "依据为既有结构化资料；转换完成不认证它等同于原始策划案，仍需核对业务含义。", "",
             "## 1. 业务背景与范围", "", "保留本次导入的要求；其他范围原资料未明确。", "",
             "## 2. 功能需求", ""]
    for req in data["requirements"]:
        lines += [f"### {md_literal(req['title'])}", f"**{req['id']}**", f"描述：{md_literal(req['description'])}",
                  f"类型：{req['kind']}；状态：{req['status']}。"]
        lines.extend(f"验收标准：{ac['id']} — {md_literal(ac['text'])}" for ac in req["acceptance_criteria"])
        lines += [f"来源：{req['source']['file']}，第{req['source']['lines'][0]}–{req['source']['lines'][1]}行。", ""]
    for number, title in [(3, "业务流程"), (4, "输入与业务约束"), (5, "非功能需求"), (6, "依赖与关联")]:
        lines += [f"## {number}. {title}", "", "详见导入要求及不可变底稿；原资料未明确的部分未补写。", ""]
    lines += ["## 7. 附图与来源", "", "来源为 imports 中保留的原JSON/表格与不可变导入底稿；旧页码仅作为旧元数据保留。", ""]
    for image in data["images"]:
        lines += [f"![{image['description']}](<{image['path']}>)", image["transcription"], ""]
    lines += ["## 8. 待澄清问题", ""]
    lines.extend(f"{q['id']}；{q['priority']}；{q['status']}：{q['text']}；解答：{q['resolution']}" for q in data["questions"])
    if not data["questions"]:
        lines.append("本次没有导入明确问题；不代表原始需求没有模糊点。")
    return "\n".join(lines) + "\n"


def normalize_requirements(rows, module, code, registry, mapping, source, batch, original_data):
    data = {"schema_version": "1.0", "module": module, "code": code, "requirements": [], "questions": [], "images": []}
    missing, used = [], set()
    transcript = [f"# {module}旧资料导入底稿", "", "来源：既有结构化需求数据。旧页码和段落锚点未经原始文档核实。", ""]
    for original, location, *extra in rows:
        row = convert_row(original, "requirements", dict(mapping, fields={}) if extra else mapping)
        req_id = make_id(row.get("id"), "REQ", code, registry, used)
        body = text(row.get("detailed_description") or row.get("description"))
        original_quote = original.get("source", {}).get("quote") if isinstance(original.get("source"), dict) else None
        quote = original_quote or body
        transcript += [f"## {text(row.get('id'))} {text(row.get('title'))}", f"原记录位置：{location}", "", "原详细描述："]
        start = len(transcript) + 1
        transcript.extend(body.splitlines() or [""])
        transcript += ["", "原简述：", *text(row.get("description")).splitlines(), "", "原来源引文："]
        transcript.extend(text(original_quote).splitlines() or [""])
        end = len(transcript)
        kind = KINDS.get(re.sub(r"\s", "", text(row.get("kind"))), text(row.get("kind")))
        status = STATUSES.get(text(row.get("status")).strip(), text(row.get("status")))
        raw_criteria = row.get("acceptance_criteria", [])
        if isinstance(raw_criteria, str):
            raw_criteria = [part for part in raw_criteria.splitlines() if part.strip()]
        if not isinstance(raw_criteria, list):
            raise ValidationError(f"{location} acceptance_criteria 必须为数组或逐行文本。")
        criteria = []
        for index, old_ac in enumerate(raw_criteria, 1):
            ac_text = text(old_ac.get("text")) if isinstance(old_ac, dict) else text(old_ac)
            legacy_ac = text(old_ac.get("id")) if isinstance(old_ac, dict) and old_ac.get("id") else f"{text(row.get('id'))}::criterion:{index}"
            new_ac = registry.setdefault("AC", {}).get(legacy_ac, f"AC_{code}_{req_id.rsplit('_', 1)[-1]}_{index:02d}" if req_id else "")
            if legacy_ac in registry["AC"] and registry["AC"][legacy_ac] != new_ac:
                raise ValidationError(f"AC映射冲突：{legacy_ac}。")
            registry["AC"][legacy_ac] = new_ac
            criteria.append({"id": new_ac, "text": ac_text})
        assets, image_problems = legacy_assets(extra[0] if extra else original, source, batch, mapping)
        # Canonical REQ may register images at the top level instead of per record.
        source_images = original.get("source", {}).get("images", []) if isinstance(original.get("source"), dict) else []
        if source_images and isinstance(original_data, dict):
            descriptors = {item.get("path"): item for item in original_data.get("images", []) if isinstance(item, dict)}
            canonical_assets, problems = legacy_assets({"related_images": [descriptors.get(name, {"path": name}) for name in source_images]}, source, batch, mapping)
            assets.extend(canonical_assets)
            image_problems.extend(problems)
        for image in assets:
            if not image["description"] or not image["transcription"]:
                image_problems.append(f"图片 {image['path']} 缺少描述或转录，需核对补充")
            if image["path"] not in {item["path"] for item in data["images"]}:
                data["images"].append(image)
        req = {"id": req_id, "title": text(row.get("title")), "description": text(row.get("description")),
               "kind": kind, "status": status, "acceptance_criteria": criteria,
               "source": {"file": os.path.relpath(batch / "input" / "导入底稿.md", batch.parents[2]).replace(os.sep, "/"),
                          "lines": [start, end], "quote": text(quote), "images": list(dict.fromkeys(i["path"] for i in assets))},
               "legacy": {"location": location, "record": extra[0] if extra else original}}
        problems = [key for key in ("id", "title", "description") if not req[key].strip()]
        if kind not in ("functional", "constraint", "nonfunctional"):
            problems.append("kind 缺失或类型待确认")
        if status not in ("active", "blocked", "deprecated"):
            problems.append("status 缺失或状态待确认；请明确active/blocked/deprecated")
        if not body or not quote:
            problems.append("原详细描述/来源引文缺失")
        if status == "active" and not criteria:
            problems.append("active需求缺少明确验收标准")
        if any(not item["text"].strip() for item in criteria):
            problems.append("验收标准文本缺失")
        problems.extend(image_problems)
        problems.extend(unresolved_markers(extra[0] if extra else original))
        req["missing_fields"] = problems
        if problems:
            missing.append({"id": req_id or f"row:{len(data['requirements']) + 1}", "location": location, "fields": problems})
        data["requirements"].append(req)
        transcript += ["", "原验收标准：", *[item["text"] for item in criteria], ""]
    # Already canonical question data is preserved; legacy numeric IDs need mapping.
    if isinstance(original_data, dict) and isinstance(original_data.get("questions"), list):
        data["questions"] = copy.deepcopy(original_data["questions"])
        for question in data["questions"]:
            question["affected_requirements"] = resolve_links(question.get("affected_requirements"), "REQ", registry)
    write_text(batch / "input" / "导入底稿.md", "\n".join(transcript) + "\n")
    return data, missing


def normalize_tests(rows, kind, module, code, registry, mapping, spec, source, batch):
    list_key = "testpoints" if kind == "testpoints" else "cases"
    id_key, prefix = ("测试点编号", "TP") if kind == "testpoints" else ("用例编号", "TC")
    data = {"schema_version": "1.0", "module": module, "code": code, list_key: []}
    missing, used = [], set()
    for original, location, *extra in rows:
        row = convert_row(original, kind, dict(mapping, fields={}) if extra else mapping)
        row[id_key] = make_id(row.get(id_key), prefix, code, registry, used)
        for field in canonical_fields(kind):
            if field not in row and field != "关联测试点":
                row[field] = [] if field in ("操作步骤", "预期结果", "关联需求", "验收标准", "设计方法") else ""
        for field in ("操作步骤", "预期结果"):
            if field in row:
                row[field] = steps(row[field])
        for field in ("前置条件", "检查内容", "测试数据"):
            if field in row:
                row[field] = text(row[field])
        row["关联需求"] = resolve_links(row["关联需求"], "REQ", registry)
        row["验收标准"], ac_problems = resolve_criteria(row["验收标准"], row["关联需求"], spec, registry)
        row["设计方法"] = items(row["设计方法"])
        if "关联测试点" in row:
            row["关联测试点"] = resolve_links(row["关联测试点"], "TP", registry)
        assets, image_problems = legacy_assets(extra[0] if extra else original, source, batch, mapping)
        row["import_images"] = [f"imports/{batch.name}/assets/{item['import_asset']}" for item in assets]
        problems = [field for field in canonical_fields(kind) if field != "关联测试点" and not row.get(field)
                    and field not in ("实际结果", "自动化脚本")]
        if row.get("优先级") not in ("P0", "P1", "P2", "P3"):
            problems.append("优先级无效或待确认")
        if row["场景类型"] not in SCOPES:
            problems.append("场景类型缺失或待确认（旧test_type保留在legacy，不等同场景类型）")
        if not row["设计方法"] or set(row["设计方法"]) - set(METHODS):
            problems.append("设计方法缺失或未知")
        if kind == "cases" and row["是否可自动化"] not in ("是", "否", "待定"):
            problems.append("是否可自动化缺失或待确认")
        reqs = {req["id"]: req for req in spec["requirements"]}
        for related in row["关联需求"]:
            if related not in reqs:
                problems.append(f"关联需求不存在：{related}")
            elif reqs[related]["status"] != "active":
                problems.append(f"关联需求不是active：{related}")
        problems.extend(ac_problems + image_problems)
        problems.extend(unresolved_markers(extra[0] if extra else original))
        row["legacy"] = {"location": location, "record": extra[0] if extra else original}
        row["missing_fields"] = list(dict.fromkeys(problems))
        if problems:
            missing.append({"id": row[id_key] or f"row:{len(data[list_key]) + 1}", "location": location, "fields": row["missing_fields"]})
        data[list_key].append(row)
    return data, missing


def check_candidate(data, kind, module_dir, batch):
    formal = copy.deepcopy(data)
    formal["draft"] = False
    if kind != "requirements":
        spec = validate_spec(module_dir)
        (validate_points(formal, spec) if kind == "testpoints" else validate_cases(formal, spec, module_dir=module_dir))
        if kind == "cases" and any(case.get("关联测试点") for case in formal["cases"]):
            point_file = module_dir / TARGETS["testpoints"]
            if not point_file.exists():
                raise ValidationError("用例含测试点关联，但当前模块没有已验证的测试点数据。")
            points = validate_points(read_json(point_file), spec)
            known = {point["测试点编号"]: point for point in points}
            for case in formal["cases"]:
                for point_id in case.get("关联测试点", []):
                    if point_id not in known:
                        raise ValidationError(f"{case['用例编号']} 关联未知测试点 {point_id}。")
                    point = known[point_id]
                    if not set(point["关联需求"]) <= set(case["关联需求"]) or not set(point["验收标准"]) <= set(case["验收标准"]):
                        raise ValidationError(f"{case['用例编号']} 未保留测试点 {point_id} 的REQ/AC归属。")
        return formal
    for sibling in module_dir.parent.glob("*/REQ清单.json"):
        if sibling.parent.resolve() == module_dir:
            continue
        try:
            other = read_json(sibling)
        except ValidationError:
            # 半写坏或旧格式的邻居只影响它自己，不能锁死本模块的导入。
            continue
        if isinstance(other, dict) and other.get("code") == data["code"]:
            raise ValidationError(f"code 与其他模块重复：{data['code']}。")
    manifest = module_dir.parent / "_公共" / "拆分清单.json"
    if manifest.exists():
        matches = [item for item in read_json(manifest).get("modules", []) if item.get("name") == data["module"]]
        if len(matches) != 1 or matches[0].get("code") != data["code"]:
            raise ValidationError("导入的module/code与现有拆分清单不一致；请先明确模块范围。")
    with tempfile.TemporaryDirectory(prefix="qa-import-check-") as temporary:
        stage = Path(temporary).resolve() / data["module"]
        stage.mkdir()
        staged = copy.deepcopy(formal)
        transcript = batch / "input" / "导入底稿.md"
        for req in staged["requirements"]:
            req["source"]["file"] = os.path.relpath(transcript, stage.parent).replace(os.sep, "/")
        for image in staged["images"]:
            destination = stage / image["path"]
            destination.parent.mkdir(exist_ok=True)
            shutil.copy2(batch / "assets" / image["import_asset"], destination)
        for name in ("解析来源.json", "解析核对.json"):
            existing = module_dir / name
            if existing.is_file():
                shutil.copy2(existing, stage / name)
        write_json(stage / TARGETS[kind], staged)
        write_text(stage / "需求规范.md", spec_markdown(staged))
        check_spec(stage, project_dir=module_dir.parent.parent)
    return formal


def publish(data, kind, module_dir, batch, replace, source):
    targets = [module_dir / TARGETS[kind]]
    if kind == "requirements":
        targets += [module_dir / "需求规范.md", module_dir / "解析.md", module_dir / "images"]
    if source.resolve() in {p.resolve() for p in targets}:
        raise ValidationError("--apply 不可覆盖原输入文件；请保留原件并从其副本导入。")
    existing = [path for path in targets if path.exists()]
    if existing and not replace:
        raise ValidationError("目标已存在：" + "、".join(p.name for p in existing) + "；需明确 --replace，发布前将完整备份。")
    if kind == "cases":
        stale_names = DERIVED[9:]
    else:
        stale_names = DERIVED
    stale = [module_dir / name for name in stale_names if (module_dir / name).exists()]
    if kind == "requirements":
        stale += [module_dir / name for name in (TARGETS["testpoints"], TARGETS["cases"]) if (module_dir / name).exists()]
    elif kind == "testpoints" and (module_dir / TARGETS["cases"]).exists():
        stale.append(module_dir / TARGETS["cases"])
    backup = batch / "备份"
    backup.mkdir()
    touched = list(dict.fromkeys(existing + stale))
    for path in touched:
        destination = backup / path.name
        if path.is_dir():
            shutil.copytree(path, destination)
        else:
            shutil.copy2(path, destination)
    try:
        for path in stale:
            path.unlink()
        if kind == "requirements":
            image_dir = module_dir / "images"
            if image_dir.exists():
                # Exact module-owned image directory was backed up above.
                shutil.rmtree(image_dir)
            image_dir.mkdir()
            for image in data["images"]:
                shutil.copy2(batch / "assets" / image["import_asset"], module_dir / image["path"])
            write_text(module_dir / "需求规范.md", spec_markdown(data))
            transcript = (batch / "input" / "导入底稿.md").read_text(encoding="utf-8")
            view = "# 旧资料导入解析\n\n原件及精确原文保留在本模块的 imports/ 导入批次中。以下仅为安全显示副本。\n\n" + md_literal(transcript)
            for image in data["images"]:
                view += f"\n![{image['description']}](<{image['path']}>)\n{image['transcription']}\n"
            write_text(module_dir / "解析.md", view)
        data["import_status"] = "applied"
        write_json(module_dir / TARGETS[kind], data)
        if kind == "requirements":
            validate_spec(module_dir)
    except Exception:
        for path in targets:
            if path.exists():
                shutil.rmtree(path) if path.is_dir() else path.unlink()
        for path in touched:
            saved = backup / path.name
            shutil.copytree(saved, path) if saved.is_dir() else shutil.copy2(saved, path)
        raise
    return [path.name for path in touched], [path.name for path in stale]


def run(args):
    source, module_dir = args.input.resolve(), args.module_dir.resolve()
    if not source.is_file():
        raise ValidationError("原输入文件不存在或不是普通文件。")
    mapping = read_json(args.mapping) if args.mapping else {}
    if not isinstance(mapping, dict):
        raise ValidationError("mapping顶层必须为对象。")
    module = mapping.get("module", module_dir.name)
    if not isinstance(module, str) or not module or module in (".", "..") or "/" in module or "\\" in module:
        raise ValidationError("module必须为单个模块名。")
    current = read_json(module_dir / TARGETS["requirements"]) if (module_dir / TARGETS["requirements"]).is_file() else None
    code = args.code or mapping.get("code") or (current.get("code") if current else None)
    original_data = read_json(source) if source.suffix.lower() == ".json" else None
    if isinstance(original_data, dict) and "schema_version" in original_data and original_data["schema_version"] != "1.0":
        raise ValidationError(f"不支持 schema_version={original_data['schema_version']}；请明确转换为1.0，避免错误解释新版字段。")
    if code is None and isinstance(original_data, dict):
        code = original_data.get("code")
    if not isinstance(code, str) or not re.fullmatch(r"[A-Z]{2,4}", code):
        raise ValidationError("请提供 --code（2–4位大写字母）；缺少模块代码时不会随机编造。")
    if current and (current.get("code") != code or current.get("module") != module):
        raise ValidationError("导入module/code与当前REQ不一致，不能改同一稳定ID的模块指代。")
    if args.kind != "requirements":
        if current and current.get("draft") is True:
            raise ValidationError("当前REQ清单是草稿，需完成核对并发布正式需求后才能导入测试点/用例。")
        spec = validate_spec(module_dir)
    else:
        spec = None
    imports = module_dir / "imports"
    imports.mkdir(parents=True, exist_ok=True)
    registry_file = imports / "编号登记.json"
    registry = read_json(registry_file) if registry_file.exists() else {"schema_version": "1.0", "module": module, "code": code}
    if registry.get("module") != module or registry.get("code") != code:
        raise ValidationError("编号登记的module/code不一致，请核对目标模块。")
    ids = mapping.get("ids", {})
    if not isinstance(ids, dict):
        raise ValidationError("mapping.ids 必须为REQ/AC/TP/TC对象。")
    for group, values in ids.items():
        if group not in ("REQ", "AC", "TP", "TC") or not isinstance(values, dict):
            raise ValidationError("mapping.ids 只支持REQ/AC/TP/TC → 旧编号到新编号的对象。")
        known = registry.setdefault(group, {})
        for old, new in values.items():
            if old in known and known[old] != new:
                raise ValidationError(f"既有稳定编号不能重新指代：{old} 已是 {known[old]}。")
            known[old] = new
    # Reserve canonical IDs in the current module before allocating imported ones.
    for group, filename, list_key, id_key in (("REQ", TARGETS["requirements"], "requirements", "id"),
                                              ("TP", TARGETS["testpoints"], "testpoints", "测试点编号"),
                                              ("TC", TARGETS["cases"], "cases", "用例编号")):
        path = module_dir / filename
        if path.exists():
            for row in read_json(path).get(list_key, []):
                registry.setdefault(group, {}).setdefault(row[id_key], row[id_key])
    batch_name = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    batch = imports / batch_name
    batch.mkdir()
    (batch / "input").mkdir()
    shutil.copy2(source, batch / "input" / source.name)
    if args.mapping:
        shutil.copy2(args.mapping, batch / "mapping.json")
    try:
        rows = json_rows(original_data, args.kind, module) if original_data is not None else table_rows(source, args.kind, args.sheet, args.header_row, mapping)
        if args.kind == "requirements":
            data, missing = normalize_requirements(rows, module, code, registry, mapping, source, batch, original_data)
        else:
            data, missing = normalize_tests(rows, args.kind, module, code, registry, mapping, spec, source, batch)
        if isinstance(original_data, dict):
            marker_problems = unresolved_markers(original_data)
            if marker_problems:
                missing.append({"id": "source", "location": "$", "fields": marker_problems})
        data.update(draft=True, import_status="incomplete" if missing else "ready_for_validation", missing_fields=missing,
                    imported_source={"file": source.name, "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                                     "kind": args.kind, "basis": "existing_structured_material"}, import_batch=batch_name)
        errors = []
        formal = None
        if not missing:
            try:
                formal = check_candidate(data, args.kind, module_dir, batch)
            except (ValidationError, ValueError, OSError) as exc:
                errors.append(str(exc))
                data["import_status"] = "validation_failed"
        data["validation_errors"] = errors
        write_json(batch / TARGETS[args.kind], data)
        write_json(batch / "编号映射.json", registry)
        # IDs are stable across batches even while unresolved drafts are retained.
        write_json(registry_file, registry)
        backups, stale = [], []
        applied = False
        if args.apply and formal is not None:
            backups, stale = publish(formal, args.kind, module_dir, batch, args.replace, source)
            applied = True
        lines = ["# 既有资料导入报告", "", f"类型：{args.kind}；模块：{module}；条目数：{len(rows)}。", "",
                 f"状态：{'正式数据已发布' if applied else data['import_status']}。规范化草稿始终 draft=true。", "",
                 "编号/字段转换仅保留既有内容；未从常识补写业务规则。依据为旧结构化资料，不认证它与原始策划案一致。",
                 "REQ关联不等于AC覆盖；只有精确唯一文本匹配或明确编号映射才绑定AC。", "",
                 "## 待补字段与校验问题", ""]
        for item in missing:
            lines.append(f"- {item['id']}（{item['location']}）：" + "；".join(item["fields"]))
        lines.extend(f"- {error}" for error in errors)
        if not missing and not errors:
            lines.append("结构与证据/关联校验通过；不代表业务已确认、覆盖充分或软件测试通过。")
        if backups:
            lines += ["", "## 备份与过期结果", "", "已备份：" + "、".join(backups), "需重新生成：" + "、".join(stale)]
        write_text(batch / "导入报告.md", "\n".join(lines) + "\n")
        if args.apply and not applied:
            raise ValidationError(f"资料尚未完整验证，不能 --apply；请查看 {batch / '导入报告.md'}。")
        return batch, applied
    except Exception as exc:
        if not (batch / "导入报告.md").exists():
            write_text(batch / "导入报告.md", f"# 导入未完成\n\n具体原因：{exc}\n\n原输入副本已保留，未发布正式数据。\n")
        raise ValidationError(str(exc)) from exc


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--module-dir", required=True, type=Path)
    parser.add_argument("--kind", required=True, choices=list(TARGETS))
    parser.add_argument("--code")
    parser.add_argument("--sheet")
    parser.add_argument("--header-row", type=int, default=1)
    parser.add_argument("--mapping", type=Path)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--replace", action="store_true")
    args = parser.parse_args()
    try:
        batch, applied = run(args)
    except Exception as exc:
        print(f"导入失败：{exc}", file=sys.stderr)
        return 1
    print(f"{'正式数据已发布；' if applied else '规范化草稿已保存；'}导入批次：{batch}")
    return 0


if __name__ == "__main__":
    sys.exit(main())

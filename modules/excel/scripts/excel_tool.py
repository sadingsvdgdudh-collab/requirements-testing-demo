#!/usr/bin/env python3
"""Read/analyze workbooks and author scoped changes with Artifact Tool only."""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
import xml.etree.ElementTree as ET


class ExcelError(ValueError):
    pass


ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}
SUPPORTED_OPS = {"set_values", "set_formulas", "format", "append_rows", "add_sheet", "add_table", "add_chart"}
MAX_CELLS = 250_000


def load_json(path: Path) -> dict:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, UnicodeError) as exc:
        raise ExcelError(f"JSON格式错误：{path.name}：{exc}") from exc
    if not isinstance(value, dict):
        raise ExcelError("Excel spec必须是JSON对象。")
    return value


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def library():
    try:
        import openpyxl
        return openpyxl
    except ImportError as exc:
        raise ExcelError("读取/验证XLSX需要openpyxl。请使用工作区提供的Python运行时。") from exc


def preflight(path: Path, writing: bool = False) -> list[str]:
    """Reject unsupported native parts BEFORE an import/export can discard them."""
    suffix = path.suffix.lower()
    if suffix != ".xlsx":
        raise ExcelError(f"此命令只接收.xlsx；{suffix or '无扩展名'}请用原应用另存副本。宏工作簿不得改后缀规避检查。")
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
            if "xl/workbook.xml" not in names:
                raise ExcelError("文件不是有效的XLSX工作簿。")
            unsupported = []
            for prefix, label in (
                ("xl/vba", "VBA宏"), ("xl/pivot", "透视表或透视缓存"),
                ("xl/slicer", "切片器"), ("xl/externalLinks/", "外部工作簿链接"),
                ("xl/embeddings/", "嵌入对象"), ("xl/ctrlProps/", "控件"),
                ("xl/activeX/", "ActiveX"), ("xl/model/", "数据模型"),
                ("xl/queryTables/", "查询表"), ("customXml/", "自定义XML"),
                ("_xmlsignatures/", "数字签名"), ("xl/threadedComments/", "线程批注"),
            ):
                if any(name.startswith(prefix) for name in names):
                    unsupported.append(label)
            if "xl/connections.xml" in names:
                unsupported.append("外部数据连接")
            for name in names:
                if name.startswith("xl/") and name.endswith(".xml"):
                    root = ET.fromstring(archive.read(name))
                    for elem in root.iter():
                        tag = elem.tag.rsplit("}", 1)[-1]
                        if tag == "extLst":
                            unsupported.append(f"扩展特性（{name}）")
                        if tag == "f" and elem.get("t") in {"array", "dataTable"}:
                            unsupported.append("数组公式或What-If Data Table")
                        if tag in {"oleObjects", "legacyDrawingHF", "controls"}:
                            unsupported.append(f"复杂对象（{tag}）")
                        if tag in {"sheetProtection", "workbookProtection"}:
                            flags = ("sheet",) if tag == "sheetProtection" else ("lockStructure", "lockWindows", "lockRevision")
                            enabled = any(elem.get(flag, "false").strip().lower() not in {"", "0", "false"} for flag in flags)
                            credentials = any(value and key.lower().endswith(("password", "algorithmname", "hashvalue", "saltvalue", "spincount"))
                                              for key, value in elem.attrib.items())
                            if enabled or credentials:
                                unsupported.append("保护设置")
                        if tag == "drawing" and any(n.startswith("xl/media/") for n in names):
                            unsupported.append("含图片的原模板绘图")
                        if name.startswith("xl/drawings/") and tag in {"sp", "cxnSp", "grpSp", "pic"}:
                            unsupported.append("形状/连接线/组合或图片绘图")
                        if tag == "r" and (name == "xl/sharedStrings.xml" or name.startswith("xl/worksheets/")):
                            unsupported.append("单元格富文本")
                        if tag == "calcPr" and elem.get("iterate") in {"1", "true"}:
                            unsupported.append("迭代计算")
                        if tag == "workbookPr" and elem.get("date1904") in {"1", "true"}:
                            unsupported.append("1904日期系统")
            if any(name.startswith("xl/drawings/vml") for name in names):
                unsupported.append("VML绘图或旧批注布局")
            unsupported = sorted(set(unsupported))
            if unsupported and writing:
                raise ExcelError("拒绝编辑/重算，当前作者不能保证保留：" + "、".join(unsupported) + "。原文件未修改。")
            return unsupported
    except (zipfile.BadZipFile, ET.ParseError) as exc:
        raise ExcelError("无法读取XLSX：文件可能已加密、损坏或并非真实XLSX。") from exc


def json_value(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    return value


def xml_signature(obj):
    if obj is None:
        return None
    if hasattr(obj, "to_tree"):
        return ET.tostring(obj.to_tree(), encoding="unicode")
    return str(obj)


def style_signature(cell):
    return tuple(xml_signature(getattr(cell, key)) for key in ("font", "fill", "border", "alignment", "protection")) + (cell.number_format,)


def chart_signature(chart):
    root = chart._write()
    # Source edits are supposed to refresh cached plot values. Keep bindings and
    # all native chart/series styles, but exclude recalculated caches from the
    # preservation comparison. Bindings and source values are checked separately.
    for parent in root.iter():
        for child in list(parent):
            if child.tag.rsplit("}", 1)[-1] in {"numCache", "strCache"}:
                parent.remove(child)
    return ET.tostring(root, encoding="unicode")


def chart_rebindings(path: Path):
    """Describe supported contiguous bindings using the original readable file."""
    wb = library().load_workbook(path, data_only=False)
    items = []
    ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart", "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    formula_pattern = re.compile(r"^(?:'((?:[^']|'')+)'|([^'!]+))!\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)$")
    for ws in wb.worksheets:
        for index, chart in enumerate(ws._charts):
            if chart.__class__.__name__ not in {"BarChart", "LineChart", "AreaChart", "PieChart", "DoughnutChart", "RadarChart"}:
                raise ExcelError("当前不保证原模板此类图表的重算/保真，请在原应用处理该图表。")
            root = ET.fromstring(ET.tostring(chart._write()))
            series = root.findall(".//c:ser", ns)
            refs, styles = [], []
            for serial, ser in enumerate(series):
                category = ser.find("c:cat//c:f", ns)
                value = ser.find("c:val//c:f", ns)
                if category is None or value is None:
                    raise ExcelError("原图表没有标准单列范围绑定，拒绝把静态或复杂图表降级重算。")
                for node in ((category, value) if serial == 0 else (value,)):
                    match = formula_pattern.fullmatch(node.text or "")
                    if not match:
                        raise ExcelError("原图表范围不是可验证的本地A1单列范围。")
                    source_sheet = (match[1] or match[2]).replace("''", "'")
                    address = f"{match[3]}{match[4]}:{match[5]}{match[6]}"
                    bounds = library().utils.range_boundaries(address)
                    if bounds[0] != bounds[2]:
                        raise ExcelError("原图表为复杂二维绑定，当前不能保证保真。")
                    refs.append((source_sheet, *bounds))
                color = ser.find("c:spPr/a:solidFill/a:srgbClr", ns)
                line = ser.find("c:spPr/a:ln", ns)
                line_color = line.find("a:solidFill/a:srgbClr", ns) if line is not None else None
                dash = line.find("a:prstDash", ns) if line is not None else None
                styles.append({"fill": "#" + color.get("val") if color is not None else None,
                               "line": {"fill": "#" + line_color.get("val"), "style": dash.get("val", "solid") if dash is not None else "solid", "width": int(line.get("w", "19050")) / 9525} if line_color is not None else None})
            if not refs or any(ref[0] != refs[0][0] or ref[2] != refs[0][2] or ref[4] != refs[0][4] for ref in refs):
                raise ExcelError("原图表类别/数据范围不对齐，当前不能保证重算。")
            if [ref[1] for ref in refs] != list(range(refs[0][1], refs[0][1] + len(refs))) or refs[0][2] < 2:
                raise ExcelError("原图表必须为包含上一行标题的连续类别/数据列；非相邻或无标题原图暂不能保真重算。")
            source_name, first_col, first_row, _, last_row = refs[0]
            last_col = refs[-1][1]
            if source_name not in wb.sheetnames or any(wb[source_name].cell(first_row - 1, col).value is None for col in range(first_col, last_col + 1)):
                raise ExcelError("原图表缺少可验证系列标题，不能盲目重绑。")
            address = f"{library().utils.get_column_letter(first_col)}{first_row - 1}:{library().utils.get_column_letter(last_col)}{last_row}"
            items.append({"sheet": ws.title, "index": index, "source_sheet": source_name, "range": address, "styles": styles})
    wb.close()
    return items


def workbook_signature(path: Path) -> dict:
    op = library()
    wb = op.load_workbook(path, data_only=False)
    result = {"order": wb.sheetnames, "sheets": {}, "names": {name: xml_signature(value) for name, value in wb.defined_names.items()}}
    for ws in wb.worksheets:
        if ws.max_row * ws.max_column > MAX_CELLS:
            raise ExcelError(f"{ws.title}有效范围超过{MAX_CELLS}格，拒绝不完整的保真验证，请拆分文件。")
        cells = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None or cell.has_style:
                    cells[cell.coordinate] = {"value": json_value(cell.value), "type": cell.data_type,
                                              "style": style_signature(cell), "hyperlink": (cell.hyperlink.target, cell.hyperlink.location, cell.hyperlink.tooltip, cell.hyperlink.display) if cell.hyperlink else None,
                                              "comment": (cell.comment.text, cell.comment.author) if cell.comment else None}
        result["sheets"][ws.title] = {
            "cells": cells, "freeze": ws.freeze_panes, "merges": sorted(str(r) for r in ws.merged_cells.ranges),
            "filter": xml_signature(ws.auto_filter), "tables": {t.name: xml_signature(t) for t in ws.tables.values()},
            "table_refs": {t.name: t.ref for t in ws.tables.values()},
            "rows": {str(k): xml_signature(v) for k, v in ws.row_dimensions.items()},
            "columns": {k: xml_signature(v) for k, v in ws.column_dimensions.items()},
            "charts": [chart_signature(c) for c in ws._charts],
            "validation": xml_signature(ws.data_validations),
            "conditional": [(str(k), [xml_signature(r) for r in rules]) for k, rules in ws.conditional_formatting._cf_rules.items()],
            "sheet_state": ws.sheet_state, "gridlines": ws.sheet_view.showGridLines,
            "print_options": xml_signature(ws.print_options), "page_setup": xml_signature(ws.page_setup),
            "page_margins": xml_signature(ws.page_margins), "print_area": str(ws.print_area),
            "print_title_rows": ws.print_title_rows, "print_title_cols": ws.print_title_cols,
        }
    wb.close()
    return result


def range_cells(address: str) -> set[str]:
    op = library()
    try:
        left, top, right, bottom = op.utils.range_boundaries(address)
    except (ValueError, TypeError) as exc:
        raise ExcelError(f"无效范围：{address}。必须为有界A1范围，不能含工作表名。") from exc
    if None in (left, top, right, bottom) or left < 1 or top < 1 or right > 16384 or bottom > 1048576 or right < left or bottom < top:
        raise ExcelError(f"无效或无界范围：{address}")
    if (right - left + 1) * (bottom - top + 1) > MAX_CELLS:
        raise ExcelError("单项操作范围过大，请拆分为有界范围。")
    return {f"{op.utils.get_column_letter(col)}{row}" for row in range(top, bottom + 1) for col in range(left, right + 1)}


def validate_matrix(address, data, formulas=False):
    op = library()
    range_cells(address)
    left, top, right, bottom = op.utils.range_boundaries(address)
    if not isinstance(data, list) or len(data) != bottom - top + 1 or any(not isinstance(row, list) or len(row) != right - left + 1 for row in data):
        raise ExcelError(f"{address}的data矩阵尺寸必须与范围完全一致。")
    for row in data:
        for value in row:
            if formulas:
                if value is not None and (not isinstance(value, str) or not value.startswith("=")):
                    raise ExcelError("公式需明确放在formulas/set_formulas，且以=开头；null表示清除。")
                if value and any(token in value.upper() for token in ("LET(", "LAMBDA(", "MAP(", "REDUCE(", "TABLE(", "WEBSERVICE(", "HYPERLINK(")):
                    raise ExcelError("新公式包含未支持的数组/网络/高级函数，请使用可验证的标量公式。")
                if value and ("[" in value and "]" in value) and "." in value.split("]")[0]:
                    raise ExcelError("不接受新建外部工作簿公式引用。")
            elif isinstance(value, dict):
                if set(value) != {"type", "value"} or value["type"] not in {"date", "datetime"}:
                    raise ExcelError("类型对象只支持{type:date|datetime,value:ISO字符串}。数值/百分比直接使用JSON数值。")
                try:
                    dt.datetime.fromisoformat(value["value"]) if value["type"] == "datetime" else dt.date.fromisoformat(value["value"])
                except (ValueError, TypeError) as exc:
                    raise ExcelError("日期必须为真实ISO日期，不接受显示字符串或缺失年份。") from exc
            elif value is not None and not isinstance(value, (str, int, float, bool)):
                raise ExcelError("单元格只能为字符串、数值、布尔、null或明确日期对象。")
            elif isinstance(value, float) and not math.isfinite(value):
                raise ExcelError("不接受NaN/Infinity数值。")


def validate_sheet(sheet: dict):
    if not isinstance(sheet, dict) or not isinstance(sheet.get("name"), str) or not sheet["name"].strip():
        raise ExcelError("每个sheet必须有非空name。")
    if len(sheet["name"]) > 31 or any(char in sheet["name"] for char in "[]:*?/\\"):
        raise ExcelError("工作表名最长31字符且不能含[]:*?/\\。")
    allowed = {"name", "values", "formulas", "formats", "column_widths", "row_heights", "freeze_rows", "freeze_columns", "tables", "charts", "show_gridlines"}
    if set(sheet) - allowed:
        raise ExcelError("sheet包含未支持的字段：" + ",".join(sorted(set(sheet) - allowed)))
    for key in ("freeze_rows", "freeze_columns"):
        if key in sheet and (not isinstance(sheet[key], int) or isinstance(sheet[key], bool) or sheet[key] < 0):
            raise ExcelError(f"{key}必须为非负整数。")
    for key in ("values", "formulas"):
        for block in sheet.get(key, []):
            validate_matrix(block.get("range"), block.get("data"), formulas=key == "formulas")
    for block in sheet.get("formats", []):
        range_cells(block.get("range"))
        validate_format(block.get("format"))
    for chart in sheet.get("charts", []):
        validate_chart(chart)
    seen_tables = {}
    for table in sheet.get("tables", []):
        validate_table(table, seen_tables)
        seen_tables[table["name"]] = table["range"]
    for key in ("column_widths", "row_heights"):
        for block in sheet.get(key, []):
            range_cells(block.get("range"))
            if not isinstance(block.get("width" if key == "column_widths" else "height"), (int, float)):
                raise ExcelError(f"{key}必须有数值width/height。")


def validate_format(value):
    allowed = {"fill", "font", "numberFormat", "borders", "horizontalAlignment", "verticalAlignment", "wrapText", "columnWidth", "rowHeight"}
    if not isinstance(value, dict) or set(value) - allowed:
        raise ExcelError("format含未知属性；支持fill/font/numberFormat/borders/对齐/wrapText/columnWidth/rowHeight。")


def validate_table(table, refs):
    cells = range_cells(table.get("range"))
    name = table.get("name")
    if not isinstance(name, str) or not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", name):
        raise ExcelError("table.name需明确唯一的英文字母/下划线开头名称，不能有空格。")
    if name in refs:
        raise ExcelError(f"表格名已存在：{name}")
    if any(cells & range_cells(address) for address in refs.values()):
        raise ExcelError("新表格与原表格重叠，拒绝创建。")


def validate_chart(value):
    if not isinstance(value, dict) or value.get("type") not in {"bar", "line", "area", "pie", "doughnut", "scatter", "radar"}:
        raise ExcelError("此版图表支持bar/line/area/pie/doughnut/scatter/radar，未支持类型不会替换为静态图片。")
    ranges = value.get("ranges") or [value.get("range")]
    for address in ranges:
        range_cells(address)
    position = value.get("position")
    if not isinstance(position, list) or len(position) != 2:
        raise ExcelError("chart.position必须给两个单元格，例[H2,N16]。")
    for address in position:
        if len(range_cells(address)) != 1:
            raise ExcelError("chart.position必须为两个单元格。")
    if not value.get("title"):
        raise ExcelError("图表必须有说明内容的title。")


def normalize_edit(spec, source):
    before = workbook_signature(source)
    names = set(before["order"])
    refs = {name: dict(sheet["table_refs"]) for name, sheet in before["sheets"].items()}
    for item in spec.get("operations", []):
        kind = item.get("op")
        if kind not in SUPPORTED_OPS:
            raise ExcelError(f"未知操作：{kind}")
        if kind == "add_sheet":
            validate_sheet(item.get("sheet"))
            name = item["sheet"]["name"]
            if name in names:
                raise ExcelError(f"工作表已存在：{name}")
            names.add(name)
            refs[name] = {t["name"]: t["range"] for t in item["sheet"].get("tables", [])}
            continue
        sheet_name = item.get("sheet")
        if sheet_name not in names:
            raise ExcelError(f"不存在工作表：{sheet_name}")
        if kind in {"set_values", "set_formulas"}:
            validate_matrix(item.get("range"), item.get("data"), formulas=kind == "set_formulas")
        elif kind == "format":
            range_cells(item.get("range"))
            validate_format(item.get("format"))
        elif kind == "add_chart":
            validate_chart(item.get("chart"))
        elif kind == "add_table":
            validate_table(item.get("table", {}), refs[sheet_name])
            refs[sheet_name][item["table"]["name"]] = item["table"]["range"]
        elif kind == "append_rows":
            data = item.get("data")
            if not isinstance(data, list) or not data or not isinstance(data[0], list) or not data[0]:
                raise ExcelError("append_rows需非空矩形data。")
            if not isinstance(item.get("start_row"), int) or item["start_row"] < 1:
                raise ExcelError("append_rows必须明确start_row，避免猜测模板尾部。")
            start_col = item.get("start_column", 1)
            if not isinstance(start_col, int) or start_col < 1:
                raise ExcelError("start_column必须为正整数。")
            op = library()
            address = f"{op.utils.get_column_letter(start_col)}{item['start_row']}:{op.utils.get_column_letter(start_col + len(data[0]) - 1)}{item['start_row'] + len(data) - 1}"
            occupied = before["sheets"].get(sheet_name, {}).get("cells", {})
            if any(occupied.get(cell, {}).get("value") is not None for cell in range_cells(address)):
                raise ExcelError("append_rows范围已有内容，请用明确set_values编辑或选择真正空白区。")
            item["range"] = address
            validate_matrix(address, data)
    return before


def verify_preservation(before: dict, after_path: Path, spec: dict):
    after = workbook_signature(after_path)
    expected_order = before["order"] + [item["sheet"]["name"] for item in spec.get("operations", []) if item["op"] == "add_sheet"]
    differences = []
    if after["order"] != expected_order:
        differences.append("工作表顺序/名称")
    if before["names"] != after["names"]:
        differences.append("已定义名称")
    for name, original in before["sheets"].items():
        current = after["sheets"].get(name)
        if current is None:
            differences.append(f"{name}丢失")
            continue
        content_allow, style_allow = set(), set()
        allow_tables, allow_charts = False, False
        row_allow, col_allow = set(), set()
        for item in spec.get("operations", []):
            if item.get("sheet") != name:
                continue
            if item["op"] in {"set_values", "set_formulas", "append_rows"}:
                content_allow.update(range_cells(item["range"]))
            if item["op"] == "format":
                style_allow.update(range_cells(item["range"]))
                left, top, right, bottom = library().utils.range_boundaries(item["range"])
                if "rowHeight" in item["format"]:
                    row_allow.update(str(r) for r in range(top, bottom + 1))
                if "columnWidth" in item["format"]:
                    col_allow.update(library().utils.get_column_letter(c) for c in range(left, right + 1))
            allow_tables |= item["op"] == "add_table"
            allow_charts |= item["op"] == "add_chart"
        for cell in set(original["cells"]) | set(current["cells"]):
            old = original["cells"].get(cell, {"value": None, "type": "n", "style": None, "hyperlink": None, "comment": None})
            new = current["cells"].get(cell, {"value": None, "type": "n", "style": None, "hyperlink": None, "comment": None})
            if cell not in content_allow and (old["value"], old["type"]) != (new["value"], new["type"]):
                differences.append(f"{name}!{cell}未请求内容")
            if cell not in style_allow and cell in original["cells"] and old["style"] != new["style"]:
                differences.append(f"{name}!{cell}未请求样式")
            if (old["hyperlink"], old["comment"]) != (new["hyperlink"], new["comment"]):
                differences.append(f"{name}!{cell}链接/批注")
        for key in ("freeze", "merges", "filter", "validation", "conditional", "sheet_state", "gridlines", "print_options", "page_setup", "page_margins", "print_area", "print_title_rows", "print_title_cols"):
            if original[key] != current[key]:
                differences.append(f"{name}/{key}")
        for key, allowed in (("rows", row_allow), ("columns", col_allow)):
            if {k: v for k, v in original[key].items() if k not in allowed} != {k: v for k, v in current[key].items() if k not in allowed}:
                differences.append(f"{name}/{key}")
        table_changed = any(current["tables"].get(key) != value for key, value in original["tables"].items())
        if table_changed or (not allow_tables and set(original["tables"]) != set(current["tables"])):
            differences.append(f"{name}/表格")
        if original["charts"] != current["charts"][:len(original["charts"])] or (not allow_charts and len(original["charts"]) != len(current["charts"])):
            differences.append(f"{name}/图表")
    if differences:
        raise ExcelError("保真检查未通过，已拒绝交付派生文件：" + "、".join(differences[:12]) + "。原模板未修改，请缩小操作或使用原应用。")
    return {"verified": True, "scope": "单元格值/公式/样式、工作表顺序、冻结、筛选、表格、图表、行列尺寸及常见原生设置"}


def read_rows(path: Path, sheet=None, address=None):
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream, delimiter="\t" if suffix == ".tsv" else ","))
        if sum(len(row) for row in rows) > MAX_CELLS:
            raise ExcelError("文本表格超过处理上限，请拆分文件。")
        return {"sheet": path.stem, "rows": rows, "formula_errors": [], "formulas": [], "cached_values": False}
    preflight(path)
    op = library()
    wb = op.load_workbook(path, data_only=False)
    cache = op.load_workbook(path, data_only=True)
    if sheet and sheet not in wb.sheetnames:
        raise ExcelError(f"不存在工作表：{sheet}")
    name = sheet or wb.sheetnames[0]
    ws, values = wb[name], cache[name]
    if address:
        range_cells(address)
        left, top, right, bottom = op.utils.range_boundaries(address)
    else:
        left, top, right, bottom = 1, 1, ws.max_column, ws.max_row
    if (bottom - top + 1) * (right - left + 1) > MAX_CELLS:
        raise ExcelError("所选范围超过处理上限，请用--range限定。")
    rows, type_rows, formulas, errors = [], [], [], []
    for row in ws.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right):
        out, kinds = [], []
        for cell in row:
            value = values[cell.coordinate].value if cell.data_type == "f" else cell.value
            out.append(json_value(value))
            kinds.append("datetime" if isinstance(value, dt.datetime) else "date" if isinstance(value, dt.date) else "boolean" if isinstance(value, bool) else "number" if isinstance(value, (int, float)) else "blank" if value is None or value == "" else "text")
            if cell.data_type == "f":
                formulas.append({"cell": cell.coordinate, "formula": cell.value, "cached_value": json_value(value), "cache_missing": value is None})
            if cell.data_type == "e" or (isinstance(value, str) and value in ERRORS):
                errors.append({"cell": cell.coordinate, "value": value})
        rows.append(out)
        type_rows.append(kinds)
    result = {"sheet": name, "rows": rows, "type_rows": type_rows, "formula_errors": errors, "formulas": formulas, "cached_values": True,
              "cache_note": "只读缓存值，不执行重算。缓存缺失或旧缓存不能证明当前计算正确。", "sheet_names": wb.sheetnames,
              "range": f"{op.utils.get_column_letter(left)}{top}:{op.utils.get_column_letter(right)}{bottom}"}
    wb.close(); cache.close()
    return result


def inspect_file(path: Path) -> dict:
    if path.suffix.lower() in {".csv", ".tsv"}:
        data = read_rows(path)
        return {"source": path.name, "sha256": sha(path), "format": path.suffix.lower()[1:], "sheets": [{"name": data["sheet"], "rows": len(data["rows"]), "columns": max((len(r) for r in data["rows"]), default=0), "sample": data["rows"][:6]}]}
    warnings = preflight(path)
    wb = library().load_workbook(path, data_only=False)
    result = {"source": path.name, "sha256": sha(path), "format": "xlsx", "unsupported_for_edit": warnings, "sheets": []}
    for ws in wb.worksheets:
        sample = read_rows(path, ws.title, f"A1:{library().utils.get_column_letter(min(ws.max_column, 8))}{min(ws.max_row, 6)}")
        result["sheets"].append({"name": ws.title, "rows": ws.max_row, "columns": ws.max_column, "freeze_panes": ws.freeze_panes,
                                "tables": list(ws.tables), "charts": len(ws._charts), "filter": ws.auto_filter.ref,
                                "merged_ranges": [str(r) for r in ws.merged_cells.ranges], "sample": sample["rows"]})
    wb.close()
    return result


def analyze_file(path, sheet=None, address=None, header_row=1):
    data = read_rows(path, sheet, address)
    rows = data.pop("rows")
    type_rows = data.pop("type_rows", None)
    if header_row < 0 or header_row > len(rows):
        raise ExcelError("header_row必须为范围内的行号，0表示没有标题。")
    width = max((len(row) for row in rows), default=0)
    titles = rows[header_row - 1] if header_row else []
    body = rows[header_row:] if header_row else rows
    body_types = type_rows[header_row:] if type_rows and header_row else type_rows
    first_column = library().utils.range_boundaries(data["range"])[0] if data.get("range") else 1
    stats = []
    for index in range(width):
        values = [row[index] if index < len(row) else None for row in body]
        present = [v for v in values if v is not None and v != ""]
        numeric = [v for v in present if isinstance(v, (int, float)) and not isinstance(v, bool)]
        types = {}
        for row_index, value in enumerate(values):
            if value is None or value == "":
                continue
            kind = body_types[row_index][index] if body_types else "boolean" if isinstance(value, bool) else "number" if isinstance(value, (int, float)) else "text"
            types[kind] = types.get(kind, 0) + 1
        stats.append({"column": library().utils.get_column_letter(first_column + index), "name": json_value(titles[index]) if index < len(titles) else None,
                      "count": len(values), "missing": len(values) - len(present), "types": types,
                      "distinct": len({json.dumps(v, ensure_ascii=False) for v in present}), "numeric_count": len(numeric),
                      "sum": sum(numeric) if numeric else None, "min": min(numeric) if numeric else None,
                      "max": max(numeric) if numeric else None, "mean": sum(numeric) / len(numeric) if numeric else None})
    return {"source": path.name, "sha256": sha(path), **data, "row_count": len(body), "columns": stats,
            "method": "逐列描述统计；不猜测业务分组、口径或缺失值。CSV/TSV字符串不自动变数字。"}


def write_report(directory: Path, name: str, data: dict):
    directory.mkdir(parents=True, exist_ok=True)
    (directory / f"{name}.json").write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    lines = [f"# {name}", "", f"来源：{data.get('source', '派生工作簿')}", ""]
    if "columns" in data:
        lines += [f"记录数：{data['row_count']}", "", "| 列 | 标题 | 缺失 | 数值数 | 合计 | 平均 |", "|---|---|---:|---:|---:|---:|"]
        for column in data["columns"]:
            lines.append("| " + " | ".join(str(column.get(k) if column.get(k) is not None else "—").replace("|", "\\|") for k in ("column", "name", "missing", "numeric_count", "sum", "mean")) + " |")
        lines += ["", data["method"], "", f"公式错误：{len(data['formula_errors'])}；缺失缓存：{sum(f['cache_missing'] for f in data['formulas'])}。", data.get("cache_note", "")]
    else:
        lines += ["```json", json.dumps(data, ensure_ascii=False, indent=2), "```"]
    (directory / f"{name}.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def csv_payload(path: Path, spec: dict):
    data = read_rows(path)
    rows = data["rows"]
    width = max((len(row) for row in rows), default=0)
    if not width or not rows:
        raise ExcelError("CSV/TSV没有可导入的数据。")
    rows = [row + [None] * (width - len(row)) for row in rows]
    settings = spec.get("csv", {})
    for column in settings.get("columns", []):
        index, kind = column.get("index"), column.get("type")
        if not isinstance(index, int) or not 1 <= index <= width or kind not in {"text", "number", "percentage", "date", "boolean"}:
            raise ExcelError("csv.columns需1-based index和text/number/percentage/date/boolean类型。")
        for line, row in enumerate(rows[1:] if settings.get("header", True) else rows, 2 if settings.get("header", True) else 1):
            value = row[index - 1]
            if value in (None, ""):
                row[index - 1] = None
                continue
            try:
                if kind == "date":
                    dt.date.fromisoformat(value)
                    value = {"type": "date", "value": value}
                elif kind in {"number", "percentage"}:
                    value = float(value)
                    if not math.isfinite(value):
                        raise ValueError("非有限数")
                elif kind == "boolean":
                    if value.lower() not in {"true", "false"}:
                        raise ValueError("需true/false")
                    value = value.lower() == "true"
            except (ValueError, AttributeError) as exc:
                raise ExcelError(f"CSV第{line}行第{index}列不能转为{kind}：{value}") from exc
            row[index - 1] = value
    right = library().utils.get_column_letter(width)
    sheet = {"name": settings.get("sheet_name", "数据"), "values": [{"range": f"A1:{right}{len(rows)}", "data": rows}],
             "formats": settings.get("formats", []), "column_widths": settings.get("column_widths", []), "freeze_rows": 1 if settings.get("header", True) else 0}
    for column in settings.get("columns", []):
        if column["type"] == "percentage" and len(rows) > 1:
            char = library().utils.get_column_letter(column["index"])
            sheet["formats"].append({"range": f"{char}2:{char}{len(rows)}", "format": {"numberFormat": "0.0%"}})
    validate_sheet(sheet)
    return {"sheets": [sheet], "checks": spec.get("checks", [])}


def runtime_config():
    config = Path(os.environ["QA_EXCEL_RUNTIME_CONFIG"]) if os.environ.get("QA_EXCEL_RUNTIME_CONFIG") else Path(__file__).with_name("runtime.json")
    runtime = load_json(config) if config.exists() else {}
    node = os.environ.get("QA_NODE_EXECUTABLE") or runtime.get("node_executable") or shutil.which("node")
    modules = os.environ.get("QA_NODE_MODULES") or runtime.get("node_modules")
    if not node or not Path(node).is_file():
        raise ExcelError("缺少Node运行时，请使用工作区依赖并设置QA_NODE_EXECUTABLE。")
    if not modules or not Path(modules).is_dir():
        raise ExcelError("缺少工作区JS依赖，请设置QA_NODE_MODULES为加载器返回的node_modules。")
    return node, Path(modules).resolve()


def run_author(payload, output: Path, spec: dict, before=None, report=None, allow_errors=False):
    node, modules = runtime_config()
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="qa-excel-", dir=output.parent) as temp:
        folder = Path(temp)
        shutil.copy2(Path(__file__).with_name("excel_tool.mjs"), folder / "excel_tool.mjs")
        (folder / "node_modules").symlink_to(modules, target_is_directory=True)
        temporary = folder / "result.xlsx"
        payload["output"] = str(temporary)
        payload["preview_dir"] = str(Path(os.environ["QA_XLSX_PREVIEW_DIR"]).resolve()) if os.environ.get("QA_XLSX_PREVIEW_DIR") else None
        completed = subprocess.run([node, str(folder / "excel_tool.mjs")], input=json.dumps(payload, ensure_ascii=False), capture_output=True, text=True, encoding="utf-8", timeout=240)
        if completed.returncode:
            raise ExcelError((completed.stderr or completed.stdout).strip()[-2500:] or "Excel作者失败。")
        if not temporary.is_file():
            raise ExcelError("作者未生成工作簿。")
        calculation = json.loads(completed.stdout.strip().splitlines()[-1])
        saved_errors, missing = [], []
        for sheet in inspect_file(temporary)["sheets"]:
            data = read_rows(temporary, sheet["name"])
            saved_errors.extend({"sheet": sheet["name"], **error} for error in data["formula_errors"])
            missing.extend({"sheet": sheet["name"], **formula} for formula in data["formulas"] if formula["cache_missing"])
        if missing:
            raise ExcelError("重算后的公式缓存缺失，拒绝把未验证结果交付为已重算：" + str(missing[:3]))
        if saved_errors and not allow_errors:
            raise ExcelError("重算发现公式错误，未交付：" + json.dumps(saved_errors[:5], ensure_ascii=False) + "。如需保留审计用错误结果，请明确--allow-formula-errors。")
        preservation = verify_preservation(before, temporary, spec) if before else {"verified": None, "scope": "新建文件"}
        os.replace(temporary, output)
    result = {"output": str(output), "sha256": sha(output), "calculation": calculation, "formula_errors": saved_errors,
              "preservation": preservation, "engine": "Artifact Tool；未声称所有Excel客户端/高级功能均支持。"}
    if report:
        write_report(report, "Excel操作报告", result)
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description="通用Excel：只读检查/统计、创建、范围编辑、公式重算、CSV/TSV导入导出。")
    parser.add_argument("command", choices=("inspect", "analyze", "create", "edit", "recalculate", "import-csv", "export-csv"))
    parser.add_argument("--input", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--spec", type=Path)
    parser.add_argument("--report", type=Path)
    parser.add_argument("--sheet")
    parser.add_argument("--range", dest="address")
    parser.add_argument("--header-row", type=int, default=1)
    parser.add_argument("--replace-output", action="store_true", help="仅允许覆盖明确的派生目标；不允许覆盖输入原件。")
    parser.add_argument("--allow-formula-errors", action="store_true", help="明确保留公式错误作为审计输出，报告不会宣称计算无错误。")
    args = parser.parse_args()
    try:
        source = args.input.resolve() if args.input else None
        output = args.output.resolve() if args.output else None
        if source and not source.is_file():
            raise ExcelError(f"输入文件不存在：{source}")
        if args.command != "create" and source is None:
            raise ExcelError("此命令需要--input。")
        if args.command not in {"inspect", "analyze"}:
            if output is None:
                raise ExcelError("作者/导出命令需要--output，默认另存新文件。")
            if source == output:
                raise ExcelError("不能覆盖输入原件，请指定不同--output。")
            if output.exists() and not args.replace_output:
                raise ExcelError("派生目标已存在，请换名称或明确--replace-output。")
        spec = load_json(args.spec.resolve()) if args.spec else {}
        if set(spec) - {"sheets", "operations", "checks", "preview_ranges", "csv"}:
            raise ExcelError("spec包含未知顶层字段。")
        if args.command == "recalculate" and spec.get("operations"):
            raise ExcelError("recalculate只接受checks/preview_ranges，不执行编辑operations；需要编辑请使用edit。")
        for check in spec.get("checks", []):
            if len(range_cells(check.get("cell"))) != 1 or not isinstance(check.get("sheet"), str) or "value" not in check:
                raise ExcelError("checks需明确sheet、单cell和期望value。")
            tolerance = check.get("tolerance", 0)
            if not isinstance(tolerance, (int, float)) or not math.isfinite(tolerance) or tolerance < 0:
                raise ExcelError("checks.tolerance必须为非负有限数。")
        if args.command == "inspect":
            result = inspect_file(source)
        elif args.command == "analyze":
            result = analyze_file(source, args.sheet, args.address, args.header_row)
        elif args.command == "export-csv":
            if output.suffix.lower() not in {".csv", ".tsv"}:
                raise ExcelError("export-csv的output需.csv或.tsv。")
            summary = inspect_file(source)
            if len(summary["sheets"]) > 1 and not args.sheet:
                raise ExcelError("工作簿有多sheet，必须--sheet明确导出哪一张。CSV不保留公式/格式/其他sheet。")
            data = read_rows(source, args.sheet, args.address)
            if any(formula["cache_missing"] for formula in data["formulas"]):
                raise ExcelError("公式缓存缺失，请先recalculate，再导出实际值。")
            output.parent.mkdir(parents=True, exist_ok=True)
            with output.open("w", encoding="utf-8-sig", newline="") as stream:
                csv.writer(stream, delimiter="\t" if output.suffix.lower() == ".tsv" else ",").writerows(data["rows"])
            result = {"output": str(output), "rows": len(data["rows"]), "sheet": data["sheet"], "formula_errors": data["formula_errors"], "note": "导出保存值和文本，不保存公式、样式或其他sheet；原件未修改。"}
        else:
            if output.suffix.lower() != ".xlsx":
                raise ExcelError("作者output必须为.xlsx。")
            before = None
            if args.command == "import-csv":
                if source.suffix.lower() not in {".csv", ".tsv"}:
                    raise ExcelError("import-csv只接收CSV/TSV。")
                spec = csv_payload(source, spec)
                payload = {"command": "create", "spec": spec}
            elif args.command == "create":
                if source:
                    raise ExcelError("create不读模板；模板请用edit，CSV请用import-csv。")
                if not spec.get("sheets"):
                    raise ExcelError("create需--spec，包含非空sheets数组。")
                names = []
                for sheet in spec["sheets"]:
                    validate_sheet(sheet)
                    if sheet["name"] in names:
                        raise ExcelError("重复sheet名称。")
                    names.append(sheet["name"])
                payload = {"command": "create", "spec": spec}
            else:
                preflight(source, writing=True)
                if args.command == "edit" and not spec.get("operations"):
                    raise ExcelError("edit需spec.operations明确范围操作。")
                before = normalize_edit(spec, source)
                payload = {"command": args.command, "input": str(source), "spec": spec, "chart_rebindings": chart_rebindings(source)}
            result = run_author(payload, output, spec, before, args.report, args.allow_formula_errors)
        if args.report and args.command in {"inspect", "analyze", "export-csv"}:
            write_report(args.report.resolve(), "Excel分析报告" if args.command == "analyze" else "Excel检查报告", result)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (ExcelError, OSError, TypeError, AttributeError, KeyError, ValueError, subprocess.TimeoutExpired) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

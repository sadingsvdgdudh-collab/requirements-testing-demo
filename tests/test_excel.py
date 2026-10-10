"""Real Artifact Tool workbook tests; requires the bundled authoring runtime."""
import csv
import datetime as dt
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "modules/excel/scripts/excel_tool.py"
EXAMPLES = ROOT / "modules/excel/examples"
loader = importlib.util.spec_from_file_location("qa_excel_tool", SCRIPT)
excel = importlib.util.module_from_spec(loader)
loader.loader.exec_module(excel)


class ExcelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        import openpyxl
        cls.op = openpyxl
        cls.folder = Path(tempfile.mkdtemp(prefix="qa-excel-tests-"))
        cls.source = cls.folder / "source.xlsx"
        cls.edited = cls.folder / "edited.xlsx"
        cls.run_cli("create", "--spec", EXAMPLES / "create-spec.json", "--output", cls.source)
        if os.environ.get("QA_XLSX_PREVIEW_DIR"):
            preview = Path(os.environ["QA_XLSX_PREVIEW_DIR"])
            for name in ("汇总", "数据"):
                shutil.copy2(preview / f"after-{name}.png", preview / f"sample-create-{name}.png")
        cls.run_cli("edit", "--input", cls.source, "--spec", EXAMPLES / "edit-spec.json", "--output", cls.edited)
        if os.environ.get("QA_XLSX_PREVIEW_DIR"):
            shutil.copy2(preview / "after-数据.png", preview / "sample-edit-数据.png")

    @classmethod
    def tearDownClass(cls):
        if os.environ.get("QA_KEEP_EXCEL_TEST_OUTPUT"):
            print(f"Excel测试产物：{cls.folder}")
        else:
            shutil.rmtree(cls.folder)

    @classmethod
    def run_cli(cls, *args, ok=True):
        result = subprocess.run([sys.executable, str(SCRIPT), *map(str, args)], capture_output=True, text=True, encoding="utf-8", timeout=250)
        if ok and result.returncode:
            raise AssertionError(result.stderr + result.stdout)
        if not ok and result.returncode == 0:
            raise AssertionError("应拒绝操作但命令成功。" + result.stdout)
        return result

    def spec_file(self, name, value):
        path = self.folder / name
        path.write_text(json.dumps(value, ensure_ascii=False), encoding="utf-8")
        return path

    def test_01_real_multisheet_formula_dates_and_chart(self):
        wb = self.op.load_workbook(self.source, data_only=False)
        self.assertEqual(wb.sheetnames, ["汇总", "数据"])
        self.assertEqual(wb["数据"]["D5"].value, "=B5*C5")
        self.assertIsInstance(wb["数据"]["E5"].value, (dt.date, dt.datetime))
        self.assertEqual(wb["数据"]["E5"].number_format, "yyyy-mm-dd")
        self.assertEqual(wb["数据"].freeze_panes, "A5")
        self.assertIn("ProductData", wb["数据"].tables)
        self.assertEqual(len(wb["数据"]._charts), 1)
        chart = wb["数据"]._charts[0]
        self.assertIn("B5:B7", chart.ser[0].val.numRef.f.replace("$", ""))
        self.assertIn("A5:A7", (chart.ser[0].cat.strRef.f if chart.ser[0].cat.strRef else chart.ser[0].cat.numRef.f).replace("$", ""))
        wb.close()
        cached = self.op.load_workbook(self.source, data_only=True)
        self.assertEqual(cached["汇总"]["B5"].value, 80)
        self.assertEqual(cached["数据"]["D7"].value, 0)
        cached.close()

    def test_02_scoped_edit_preserves_template(self):
        before = excel.workbook_signature(self.source)
        result = excel.verify_preservation(before, self.edited, excel.load_json(EXAMPLES / "edit-spec.json"))
        self.assertTrue(result["verified"])
        data = self.op.load_workbook(self.edited, data_only=True)
        self.assertEqual(data["数据"]["D5"].value, 50)
        self.assertEqual(data["汇总"]["B5"].value, 110)
        data.close()

    def test_03_saved_chart_binding_and_changed_source(self):
        ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
        with zipfile.ZipFile(self.edited) as archive:
            parts = [name for name in archive.namelist() if "/charts/" in name and name.endswith(".xml")]
            self.assertEqual(len(parts), 1)
            root = excel.ET.fromstring(archive.read(parts[0]))
            binding = root.find(".//c:val/c:numRef/c:f", ns)
            self.assertEqual(binding.text.replace("$", ""), "'数据'!B5:B7")
        wb = self.op.load_workbook(self.edited, data_only=True)
        self.assertEqual(wb["数据"]["B5"].value, 5)
        wb.close()
        # This runtime writes empty chart caches even on newly created files.
        # Bindings + saved source are verified here; visual QA checks the chart,
        # and this test makes no claim of native Excel client recalculation.

    def test_04_inspect_and_analysis_report(self):
        report = self.folder / "analysis"
        result = self.run_cli("analyze", "--input", self.edited, "--sheet", "数据", "--range", "A4:E7", "--report", report)
        data = json.loads(result.stdout)
        self.assertEqual(data["row_count"], 3)
        self.assertEqual(data["columns"][1]["sum"], 8)
        self.assertEqual(data["columns"][3]["sum"], 110)
        self.assertEqual(data["formula_errors"], [])
        self.assertTrue((report / "Excel分析报告.md").is_file())
        inspected = json.loads(self.run_cli("inspect", "--input", self.edited).stdout)
        self.assertEqual(inspected["sheets"][1]["tables"], ["ProductData"])

    def test_05_recalculate_and_preserve(self):
        output = self.folder / "recalculated.xlsx"
        check = self.spec_file("recalc-checks.json", {"checks": [{"sheet": "汇总", "cell": "B5", "value": 110}]})
        data = json.loads(self.run_cli("recalculate", "--input", self.edited, "--output", output, "--spec", check).stdout)
        self.assertTrue(data["calculation"]["recalculated"])
        self.assertTrue(data["preservation"]["verified"])

    def test_06_csv_import_types_literal_and_missing(self):
        source = self.folder / "input.csv"
        source.write_text('编号,金额,比例,日期,说明\n001,10,0.15,2026-10-01,"=SUM(1,2)"\n002,,0,2026-10-02,普通文字\n', encoding="utf-8")
        settings = self.spec_file("csv-settings.json", {"csv": {"sheet_name": "导入", "columns": [{"index": 2, "type": "number"}, {"index": 3, "type": "percentage"}, {"index": 4, "type": "date"}]}})
        output = self.folder / "csv.xlsx"
        self.run_cli("import-csv", "--input", source, "--output", output, "--spec", settings)
        wb = self.op.load_workbook(output, data_only=False)
        self.assertEqual(wb["导入"]["A2"].value, "001")
        self.assertEqual(wb["导入"]["B2"].value, 10)
        self.assertIsNone(wb["导入"]["B3"].value)
        self.assertEqual(wb["导入"]["C2"].number_format, "0.0%")
        self.assertEqual(wb["导入"]["E2"].data_type, "s")
        self.assertEqual(wb["导入"]["E2"].value, "=SUM(1,2)")
        wb.close()

    def test_07_tsv_import_and_export(self):
        source = self.folder / "input.tsv"
        source.write_text("编号\t说明\n001\t含逗号,及中文\n002\t空值\n", encoding="utf-8")
        workbook = self.folder / "tsv.xlsx"
        self.run_cli("import-csv", "--input", source, "--output", workbook)
        output = self.folder / "export.tsv"
        self.run_cli("export-csv", "--input", workbook, "--output", output)
        with output.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream, delimiter="\t"))
        self.assertEqual(rows[1], ["001", "含逗号,及中文"])

    def test_08_original_and_output_protection(self):
        original = hashlib.sha256(self.source.read_bytes()).hexdigest()
        result = self.run_cli("edit", "--input", self.source, "--output", self.source, "--spec", EXAMPLES / "edit-spec.json", ok=False)
        self.assertIn("不能覆盖输入原件", result.stderr)
        self.run_cli("create", "--spec", EXAMPLES / "create-spec.json", "--output", self.source, ok=False)
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), original)

    def test_09_macro_and_pivot_parts_refused(self):
        for part in ("xl/vbaProject.bin", "xl/pivotTables/pivotTable1.xml"):
            source = self.folder / ("macro.xlsx" if "vba" in part else "pivot.xlsx")
            shutil.copy2(self.source, source)
            with zipfile.ZipFile(source, "a") as archive:
                archive.writestr(part, b"placeholder" if part.endswith(".bin") else b"<pivotTableDefinition/>")
            result = self.run_cli("recalculate", "--input", source, "--output", self.folder / (source.stem + "-out.xlsx"), ok=False)
            self.assertIn("拒绝编辑/重算", result.stderr)
            self.assertFalse((self.folder / (source.stem + "-out.xlsx")).exists())

    def test_10_xlsm_extension_refused(self):
        source = self.folder / "fake.xlsm"
        shutil.copy2(self.source, source)
        result = self.run_cli("recalculate", "--input", source, "--output", self.folder / "xlsm-out.xlsx", ok=False)
        self.assertIn("宏工作簿不得改后缀", result.stderr)

    def test_11_bad_matrix_and_unknown_operation(self):
        for spec in ({"sheets": [{"name": "Bad", "values": [{"range": "A1:B2", "data": [[1, 2]]}]}]},
                     {"operations": [{"op": "remove_sheet", "sheet": "数据"}]}):
            path = self.spec_file("invalid.json", spec)
            command = "create" if "sheets" in spec else "edit"
            args = [command, "--spec", path, "--output", self.folder / "bad.xlsx"]
            if command == "edit": args += ["--input", self.source]
            result = self.run_cli(*args, ok=False)
            self.assertNotIn("Traceback", result.stderr)

    def test_12_check_mismatch_refuses_output(self):
        spec = {"sheets": [{"name": "数据", "values": [{"range": "A1", "data": [[2]]}]}], "checks": [{"sheet": "数据", "cell": "A1", "value": 999}]}
        result = self.run_cli("create", "--spec", self.spec_file("bad-check.json", spec), "--output", self.folder / "bad-check.xlsx", ok=False)
        self.assertIn("计算核对失败", result.stderr)
        self.assertFalse((self.folder / "bad-check.xlsx").exists())

    def test_13_formula_errors_not_claimed_success(self):
        spec = {"sheets": [{"name": "数据", "formulas": [{"range": "A1", "data": [["=1/0"]]}]}]}
        result = self.run_cli("create", "--spec", self.spec_file("bad-formula.json", spec), "--output", self.folder / "bad-formula.xlsx", ok=False)
        self.assertIn("公式错误", result.stderr)
        self.assertFalse((self.folder / "bad-formula.xlsx").exists())

    def test_14_missing_input_stays_unavailable(self):
        spec = {"sheets": [{"name": "数据", "values": [{"range": "A1:C2", "data": [["项", "数量", "单价"], ["A", None, 10]]}],
                            "formulas": [{"range": "D2", "data": [[ '=IF(B2="","n.a.",B2*C2)' ]]}]}],
                "checks": [{"sheet": "数据", "cell": "D2", "value": "n.a."}]}
        output = self.folder / "missing.xlsx"
        self.run_cli("create", "--spec", self.spec_file("missing.json", spec), "--output", output)
        data = excel.analyze_file(output)
        self.assertEqual(data["columns"][1]["missing"], 1)
        self.assertEqual(data["columns"][3]["numeric_count"], 0)

    def test_15_append_does_not_overwrite(self):
        spec = {"operations": [{"op": "append_rows", "sheet": "数据", "start_row": 5, "data": [["overwrite"]]}]}
        result = self.run_cli("edit", "--input", self.source, "--spec", self.spec_file("overwrite.json", spec), "--output", self.folder / "overwrite.xlsx", ok=False)
        self.assertIn("已有内容", result.stderr)

    def test_16_csv_invalid_numeric_exposes_position(self):
        path = self.folder / "invalid.csv"
        path.write_text("金额\n不是数字\n", encoding="utf-8")
        spec = self.spec_file("numeric.json", {"csv": {"columns": [{"index": 1, "type": "number"}]}})
        result = self.run_cli("import-csv", "--input", path, "--spec", spec, "--output", self.folder / "numeric.xlsx", ok=False)
        self.assertIn("第2行第1列", result.stderr)

    def test_17_csv_multisheet_requires_choice(self):
        result = self.run_cli("export-csv", "--input", self.source, "--output", self.folder / "all.csv", ok=False)
        self.assertIn("必须--sheet", result.stderr)

    def test_18_unsupported_chart_not_silently_substituted(self):
        spec = {"sheets": [{"name": "数据", "charts": [{"type": "waterfall", "range": "A1:B3", "title": "瀑布图", "position": ["E1", "J15"]}]}]}
        result = self.run_cli("create", "--spec", self.spec_file("chart.json", spec), "--output", self.folder / "chart.xlsx", ok=False)
        self.assertIn("不会替换为静态图片", result.stderr)

    def test_19_append_and_new_sheet_preserve_existing(self):
        spec = {"operations": [{"op": "append_rows", "sheet": "数据", "start_row": 20, "data": [["附加说明", "保留原表"]]},
                              {"op": "add_sheet", "sheet": {"name": "补充", "values": [{"range": "A1:B2", "data": [["指标", "金额"], ["总金额", None]]}],
                                                             "formulas": [{"range": "B2", "data": [["='汇总'!B5"]]}]}}],
                "checks": [{"sheet": "补充", "cell": "B2", "value": 80}]}
        output = self.folder / "append.xlsx"
        result = json.loads(self.run_cli("edit", "--input", self.source, "--spec", self.spec_file("append.json", spec), "--output", output).stdout)
        self.assertTrue(result["preservation"]["verified"])
        wb = self.op.load_workbook(output, data_only=False)
        self.assertEqual(wb["数据"]["A20"].value, "附加说明")
        self.assertEqual(wb.sheetnames, ["汇总", "数据", "补充"])
        wb.close()

    def test_20_scoped_format_preserves_other_styles(self):
        spec = {"operations": [{"op": "format", "sheet": "数据", "range": "B5", "format": {"fill": "#FFF2CC"}}]}
        output = self.folder / "format.xlsx"
        result = json.loads(self.run_cli("edit", "--input", self.source, "--spec", self.spec_file("format.json", spec), "--output", output).stdout)
        self.assertTrue(result["preservation"]["verified"])
        before = self.op.load_workbook(self.source)
        after = self.op.load_workbook(output)
        self.assertNotEqual(excel.xml_signature(before["数据"]["B5"].fill), excel.xml_signature(after["数据"]["B5"].fill))
        self.assertEqual(excel.xml_signature(before["数据"]["B6"].fill), excel.xml_signature(after["数据"]["B6"].fill))
        before.close(); after.close()

    def test_21_recalculate_cannot_hide_edits(self):
        result = self.run_cli("recalculate", "--input", self.source, "--spec", EXAMPLES / "edit-spec.json", "--output", self.folder / "hidden-edit.xlsx", ok=False)
        self.assertIn("不执行编辑operations", result.stderr)

    def test_22_date_type_and_range_column_stats(self):
        data = excel.analyze_file(self.source, "数据", "B4:E7")
        self.assertEqual(data["columns"][0]["column"], "B")
        self.assertEqual(data["columns"][3]["types"], {"datetime": 3})

    def test_23_shapes_rich_text_and_iteration_refused(self):
        for label, part, contents in (
            ("shape", "xl/drawings/unsupported-shape.xml", "<drawing><sp/></drawing>"),
            ("rich", "xl/worksheets/unsupported-rich.xml", "<worksheet><r><t>富文本</t></r></worksheet>"),
            ("iteration", "xl/unsupported-iteration.xml", '<workbook><calcPr iterate="1"/></workbook>'),
        ):
            source = self.folder / (label + ".xlsx")
            shutil.copy2(self.source, source)
            with zipfile.ZipFile(source, "a") as archive:
                archive.writestr(part, contents)
            result = self.run_cli("recalculate", "--input", source, "--output", self.folder / (label + "-out.xlsx"), ok=False)
            self.assertIn("拒绝编辑/重算", result.stderr)
            self.assertFalse((self.folder / (label + "-out.xlsx")).exists())


if __name__ == "__main__":
    unittest.main()

"""Real document/image fixtures, provenance and the OCR approval gate."""
import base64
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest import mock
import zipfile

PACKAGE = Path(__file__).resolve().parents[1]
SCRIPTS = PACKAGE / "modules/document-parser/scripts"
sys.path.insert(0, str(SCRIPTS))
import parse_document as parser
import check_transcription as review


class ParserTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="qa-parser-")
        self.project = Path(self.temp.name)
        self.module = self.project / "output/任务奖励"

    def tearDown(self):
        self.temp.cleanup()

    def run_parser(self, name, mode="off", ok=True):
        source = PACKAGE / "examples-v12/document-inputs" / name
        before = source.read_bytes()
        result = subprocess.run([sys.executable, str(SCRIPTS / "parse_document.py"), "--input", str(source),
                                 "--project-dir", str(self.project), "--module", "任务奖励", "--ocr", mode],
                                capture_output=True, text=True, timeout=180)
        self.assertEqual(result.returncode, 0 if ok else 1, result.stdout + result.stderr)
        self.assertNotIn("Traceback", result.stderr)
        self.assertEqual(before, source.read_bytes())
        if ok:
            return json.loads((self.module / "解析来源.json").read_text(encoding="utf-8"))
        return result

    def approve(self, manifest, findings=None):
        data = {"schema_version": "1.0", "original_sha256": manifest["original"]["sha256"],
                "draft_sha256": manifest["draft_sha256"], "status": "confirmed", "reviewer": "测试核对者",
                "notes": "测试构造核对记录，仅验证门禁；不认证真实业务。", "findings": findings or []}
        (self.module / "解析核对.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        return data

    def test_docx_order_and_content(self):
        self.run_parser("任务规则.docx")
        text = (self.module / "解析.md").read_text(encoding="utf-8")
        self.assertLess(text.index("领取条件"), text.index("界面规则示例"))
        self.assertIn("任务完成3次，可领取10金币。", text)

    def test_docx_merged_table_metadata(self):
        data = self.run_parser("任务规则.docx")
        table = next(b for b in data["blocks"] if b["kind"] == "table")
        self.assertEqual(table["merges"][0]["colspan"], 3)
        self.assertIn("每日限领取1次。", table["rows"][-1])

    def test_docx_embedded_image_is_byte_exact(self):
        data = self.run_parser("任务规则.docx")
        with zipfile.ZipFile(PACKAGE / "examples-v12/document-inputs/任务规则.docx") as archive:
            original = archive.read("word/media/image1.png")
        saved = (self.module / "images" / data["images"][0]["name"]).read_bytes()
        self.assertEqual(original, saved)
        self.assertTrue(data["needs_review"])

    def test_text_pdf_real_page_and_table(self):
        data = self.run_parser("任务规则文本.pdf")
        self.assertEqual(data["blocks"][0]["location"], {"page": 1})
        table = next(b for b in data["blocks"] if b["kind"] == "table")
        self.assertEqual(table["rows"][1][0], "任务次数不足3")
        self.assertEqual(len(table["location"]["bbox"]), 4)

    def test_scan_pdf_without_ocr_is_not_approved(self):
        data = self.run_parser("任务规则扫描.pdf")
        self.assertEqual(len(data["images"]), 1)
        self.assertTrue(data["needs_review"])
        self.assertNotIn("任务完成3", (self.module / "解析.md").read_text(encoding="utf-8"))

    def test_screenshot_without_ocr_retains_original(self):
        data = self.run_parser("任务规则截图.png")
        self.assertEqual(data["original"]["sha256"], data["images"][0]["sha256"])
        self.assertTrue(data["needs_review"])

    def test_original_copy_and_draft_hashes(self):
        data = self.run_parser("任务规则文本.pdf")
        self.assertEqual(parser.digest(self.project / data["original"]["file"]), data["original"]["sha256"])
        self.assertEqual(parser.digest(self.project / data["draft"]), data["draft_sha256"])

    def test_block_ranges_refer_to_real_lines(self):
        data = self.run_parser("任务规则.docx")
        lines = (self.project / data["draft"]).read_text(encoding="utf-8").splitlines()
        for block in data["blocks"]:
            start, end = block["lines"]
            self.assertTrue(1 <= start <= end <= len(lines))
            self.assertTrue("\n".join(lines[start-1:end]).strip())

    def test_repeat_same_source_does_not_rewrite(self):
        self.run_parser("任务规则.docx")
        before = (self.module / "解析.md").read_bytes()
        self.run_parser("任务规则.docx")
        self.assertEqual(before, (self.module / "解析.md").read_bytes())

    def test_immutable_bytes_refuses_conflict(self):
        target = self.project / "evidence.txt"
        parser.immutable_bytes(target, b"first")
        with self.assertRaises(parser.ValidationError):
            parser.immutable_bytes(target, b"second")
        self.assertEqual(target.read_bytes(), b"first")

    def test_rejected_reparse_leaves_no_unregistered_images(self):
        """解析.md 冲突被拒时，第二份文档的图片不得留在 images/ 里。

        残留会让 validate_spec 的"images 磁盘与登记不一致"永久失败，
        用户只能手工删图。
        """
        first = self.run_parser("任务规则.docx")
        registered = {item["name"] for item in first["images"]}
        self.assertTrue(registered)
        from docx import Document
        from io import BytesIO
        # 1x1 PNG：与示例文档内嵌图片内容不同，因此图片名也不同。
        png = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8DwHwAFAAH/q842iQAAAABJRU5ErkJggg==")
        second = self.project / "第二份.docx"
        doc = Document()
        doc.add_paragraph("另一份文档的规则文本。")
        doc.add_picture(BytesIO(png))
        doc.save(second)
        result = subprocess.run([sys.executable, str(SCRIPTS / "parse_document.py"), "--input", str(second),
                                 "--project-dir", str(self.project), "--module", "任务奖励", "--ocr", "off"],
                                capture_output=True, text=True, timeout=180)
        self.assertEqual(result.returncode, 1, result.stdout + result.stderr)
        self.assertIn("已存在且内容不同", result.stderr)
        self.assertEqual({path.name for path in (self.module / "images").iterdir()}, registered)

    def test_unsupported_doc_fails_clearly(self):
        source = self.project / "old.doc"
        source.write_bytes(b"unsupported")
        transcript = parser.Transcription(self.project, self.module, "test", source)
        with self.assertRaisesRegex(parser.ValidationError, "docx"):
            parser.parse_input(source, transcript, "off")

    def test_markdown_with_unarchived_image_is_rejected(self):
        source = self.project / "source.md"
        source.write_text("![diagram](https://example.invalid/image.png)", encoding="utf-8")
        transcript = parser.Transcription(self.project, self.module, "test", source)
        with self.assertRaisesRegex(parser.ValidationError, "避免丢图"):
            parser.parse_input(source, transcript, "off")

    def test_nested_table_text_survives(self):
        from docx import Document
        source = self.project / "nested.docx"
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "父单元格"
        cell.add_table(rows=1, cols=1).cell(0, 0).text = "嵌套规则：奖励99金币"
        doc.sections[0].header.paragraphs[0].text = "页眉规则：版本2"
        doc.save(source)
        transcript = parser.Transcription(self.project, self.module, "test", source)
        parser.parse_docx(source, transcript, "off")
        self.assertIn("奖励99金币", "\n".join(transcript.lines))
        self.assertIn("页眉规则：版本2", "\n".join(transcript.lines))
        self.assertTrue(any("嵌套" in w for w in transcript.warnings))

    def test_word_auto_list_has_review_warning(self):
        from docx import Document
        source = self.project / "list.docx"
        doc = Document()
        doc.add_paragraph("规则编号需要对照原件", style="List Number")
        doc.save(source)
        transcript = parser.Transcription(self.project, self.module, "test", source)
        parser.parse_docx(source, transcript, "off")
        self.assertIn("规则编号需要对照原件", "\n".join(transcript.lines))
        self.assertTrue(any("自动列表编号" in w for w in transcript.warnings))

    def test_password_pdf_reports_clear_failure(self):
        from pypdf import PdfWriter
        source = self.project / "locked.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        writer.encrypt("secret")
        writer.write(source)
        transcript = parser.Transcription(self.project, self.module, "test", source)
        with self.assertRaisesRegex(parser.ValidationError, "密码"):
            parser.parse_pdf(source, transcript, "off")

    def test_unchecked_ocr_cannot_become_formal(self):
        self.run_parser("任务规则截图.png")
        with self.assertRaisesRegex(parser.ValidationError, "尚未视觉核对"):
            review.check(self.module)

    def test_matching_review_record_passes_hash_checks(self):
        data = self.run_parser("任务规则截图.png")
        self.approve(data)
        self.assertEqual(review.check(self.module)["original"], data["original"])

    def test_changed_original_rejects_review(self):
        data = self.run_parser("任务规则截图.png")
        self.approve(data)
        (self.project / data["original"]["file"]).write_bytes(b"changed")
        with self.assertRaisesRegex(parser.ValidationError, "已变"):
            review.check(self.module)

    def test_stale_review_rejected(self):
        data = self.run_parser("任务规则截图.png")
        record = self.approve(data)
        record["draft_sha256"] = "old"
        (self.module / "解析核对.json").write_text(json.dumps(record), encoding="utf-8")
        with self.assertRaisesRegex(parser.ValidationError, "过期"):
            review.check(self.module)

    def test_uncertain_transcription_blocks_only_named_requirements(self):
        data = self.run_parser("任务规则截图.png")
        record = self.approve(data, [{"status": "uncertain", "notes": "金额不清", "affected_requirements": ["REQ_TSK_001"]}])
        record["status"] = "confirmed_with_issues"
        (self.module / "解析核对.json").write_text(json.dumps(record), encoding="utf-8")
        review.check(self.module, [{"id": "REQ_TSK_001", "status": "blocked"}, {"id": "REQ_TSK_002", "status": "active"}])
        with self.assertRaisesRegex(parser.ValidationError, "blocked"):
            review.check(self.module, [{"id": "REQ_TSK_001", "status": "active"}])

    def test_manifest_path_escape_rejected(self):
        data = self.run_parser("任务规则截图.png")
        data["draft"] = "../../outside.md"
        (self.module / "解析来源.json").write_text(json.dumps(data), encoding="utf-8")
        with self.assertRaisesRegex(parser.ValidationError, "越界"):
            review.check(self.module)


@unittest.skipUnless(os.environ.get("QA_TEST_OCR") == "1", "Set QA_TEST_OCR=1 for native OCR acceptance")
class OCRTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="qa-parser-ocr-")
        cls.project = Path(cls.temp.name)
        cls.data = {}
        for label, name in (("image", "任务规则截图.png"), ("docx", "任务规则.docx"), ("scan", "任务规则扫描.pdf")):
            result = subprocess.run([sys.executable, str(SCRIPTS / "parse_document.py"), "--input", str(PACKAGE / "examples-v12/document-inputs" / name),
                                     "--project-dir", str(cls.project), "--module", label], capture_output=True, text=True, timeout=180)
            if result.returncode:
                raise AssertionError(result.stdout + result.stderr)
            cls.data[label] = json.loads((cls.project / "output" / label / "解析来源.json").read_text(encoding="utf-8"))

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def assert_rules(self, label):
        data = self.data[label]
        ocr = [b for b in data["blocks"] if b["kind"] == "ocr"]
        text = "\n".join(line["text"] for b in ocr for line in b["ocr"]["lines"])
        for rule in ("3", "10", "1", "金币", "未完成"):
            self.assertIn(rule, text)
        self.assertTrue(data["needs_review"])
        self.assertTrue(all(b["needs_review"] for b in ocr))

    def test_native_chinese_screenshot_ocr(self):
        self.assert_rules("image")

    def test_docx_embedded_image_ocr(self):
        self.assert_rules("docx")

    def test_scanned_pdf_ocr(self):
        self.assert_rules("scan")


if __name__ == "__main__":
    unittest.main(verbosity=2)
